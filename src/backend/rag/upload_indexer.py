"""Parse and chunk uploaded session files for per-session RAG indexing.

Uploads already land on disk under the file-read root (see
``src/frontend/chat/uploads.py``); this module turns a saved file into
``Document`` chunks that carry ``metadata["source"]`` = the tool-ready relative
path, so the phase-1 citation plumbing attributes answers to the file
automatically. Parsing stays dependency-light and mirrors the file tools:
stdlib zipfile/ElementTree for .docx, pypdf for .pdf, openpyxl for .xlsx,
python-pptx for .pptx, plain UTF-8 for everything text-like.
"""

from __future__ import annotations

import logging
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from langchain_core.documents import Document

from src.config import Settings

from .document_loader import split_documents, split_markdown_sections
from .document_quality import DocumentQualityConfig, filter_quality_documents
from .llamaindex_adapter import HierarchicalSplitConfig, split_documents_hierarchical

logger = logging.getLogger(__name__)

# WordprocessingML namespace; paragraphs are <w:p>, text runs are <w:t> —
# same convention as the word tool's extractor.
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DOCUMENT_PART = "word/document.xml"

# Suffixes whose text can be embedded after parsing. Mirrors the upload
# allowlist minus archives (.zip): an archive is not a document.
INDEXABLE_SUFFIXES = (
    ".txt",
    ".md",
    ".log",
    ".csv",
    ".docx",
    ".xlsx",
    ".pptx",
    ".pdf",
    ".ts",
    ".tsx",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".r",
    ".rs",
    ".go",
    ".sql",
    ".php",
    ".rb",
    ".tex",
    ".pl",
    ".hs",
    ".lua",
    ".jl",
    ".sh",
    ".bash",
    ".m",
    ".groovy",
    ".swift",
)

MAX_UPLOAD_DOC_CHARS = 2_000_000


class UploadIndexingError(Exception):
    """Raised when one uploaded file cannot be parsed or indexed."""


@dataclass(frozen=True)
class UploadChunkResult:
    """Parsed/chunked output for one uploaded file."""

    source: str
    chunks: list[Document]
    # Parent docstore additions when hierarchical chunking produced them.
    parent_docstore_additions: Any | None = None


def _to_metadata(source: str) -> dict[str, Any]:
    return {"source": source, "kind": "upload"}


def parse_upload_document(path: Path, *, source: str) -> list[Document]:
    """Parse one saved upload into Documents tagged with its relative path.

    ``source`` is the tool-ready relative path used both for citation display
    and for idempotent re-indexing (chunks are deleted by source first).
    """

    suffix = path.suffix.lower()
    if suffix not in INDEXABLE_SUFFIXES:
        raise UploadIndexingError(
            f"unsupported file type {suffix or '(none)'} for indexing; supported: "
            f"{', '.join(INDEXABLE_SUFFIXES)}."
        )
    if not path.is_file():
        raise UploadIndexingError(f"file not found: {path.name}.")

    if suffix == ".docx":
        texts = _extract_docx_paragraphs(path)
    elif suffix == ".pdf":
        texts = _extract_pdf_pages(path)
    elif suffix == ".xlsx":
        texts = _extract_xlsx_lines(path)
    elif suffix == ".pptx":
        texts = _extract_pptx_lines(path)
    else:
        texts = _read_text_lines(path)

    joined = "\n\n".join(texts).strip()
    if not joined:
        raise UploadIndexingError(
            f"{path.name} contains no extractable text, so it cannot be indexed."
        )
    # One Document per page/paragraph-sized unit keeps later quality filtering
    # and chunking at a granularity where a bad section does not poison the file.
    units = [unit.strip() for unit in texts if unit and unit.strip()]
    if not units:
        units = [joined]
    bounded: list[str] = []
    total = 0
    for unit in units:
        if total + len(unit) > MAX_UPLOAD_DOC_CHARS:
            logger.info(
                "Truncated %s at %d chars for indexing (limit %d)",
                path.name,
                total,
                MAX_UPLOAD_DOC_CHARS,
            )
            break
        bounded.append(unit)
        total += len(unit)
    metadata = _to_metadata(source)
    return [Document(page_content=unit, metadata=dict(metadata)) for unit in bounded]


def _extract_docx_paragraphs(path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(path) as archive, archive.open(_DOCUMENT_PART) as document:
            tree = ET.parse(document)
    except (OSError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise UploadIndexingError(f"could not read the Word document: {exc}") from exc

    paragraphs: list[str] = []
    for paragraph in tree.iter(f"{{{_W_NS}}}p"):
        texts = [node.text for node in paragraph.iter(f"{{{_W_NS}}}t") if node.text]
        line = "".join(texts).strip()
        if line:
            paragraphs.append(line)
    return paragraphs


def _extract_pdf_pages(path: Path) -> list[str]:
    try:
        import pypdf
    except ImportError as exc:
        raise UploadIndexingError("pypdf is not installed; cannot index PDF uploads.") from exc

    try:
        reader = pypdf.PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise UploadIndexingError(f"could not read the PDF: {exc}") from exc
    return [page.strip() for page in pages if page.strip()]


def _extract_xlsx_lines(path: Path) -> list[str]:
    try:
        import openpyxl
    except ImportError as exc:
        raise UploadIndexingError("openpyxl is not installed; cannot index Excel uploads.") from exc

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"# sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if value is None else str(value) for value in row]
                    if any(cell.strip() for cell in cells):
                        lines.append(" | ".join(cells))
            return lines
        finally:
            workbook.close()
    except Exception as exc:
        raise UploadIndexingError(f"could not read the spreadsheet: {exc}") from exc


def _extract_pptx_lines(path: Path) -> list[str]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise UploadIndexingError(
            "python-pptx is not installed; cannot index PowerPoint uploads."
        ) from exc

    try:
        presentation = Presentation(str(path))
        lines: list[str] = []
        for index, slide in enumerate(presentation.slides, 1):
            lines.append(f"# slide {index}")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = "\n".join(
                        paragraph.text for paragraph in shape.text_frame.paragraphs
                    ).strip()
                    if text:
                        lines.append(text)
        return lines
    except Exception as exc:
        raise UploadIndexingError(f"could not read the presentation: {exc}") from exc


def _read_text_lines(path: Path) -> list[str]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise UploadIndexingError(f"could not read the file: {exc}") from exc
    # Paragraph units, not raw lines: code files stay whole where blank-line
    # separated, prose keeps paragraph granularity.
    return [unit.strip() for unit in text.split("\n\n") if unit.strip()] or [text]


def quality_config_from_settings(settings: Settings) -> DocumentQualityConfig:
    return DocumentQualityConfig(
        enabled=settings.document_quality_filter_enabled,
        min_text_length=settings.document_quality_min_text_length,
        min_unique_terms=settings.document_quality_min_unique_terms,
        relevance_query=settings.document_quality_relevance_query,
        min_query_term_overlap=settings.document_quality_query_min_overlap,
        min_similarity=settings.document_quality_min_similarity,
        recency_bias_days=settings.document_quality_recency_bias_days,
    )


def chunk_upload_documents(
    documents: list[Document],
    *,
    settings: Settings,
    source: str,
    chat_model: Any | None = None,
    chunk_context_cache: Any | None = None,
) -> UploadChunkResult:
    """Quality-filter, then chunk upload documents like source-URL indexing.

    Mirrors the plain/hierarchical branch of
    ``ChromaRetriever._build_langchain_retriever`` so an indexed upload is
    shaped exactly like a fetched web page.
    """

    kept = filter_quality_documents(
        documents,
        quality_config_from_settings(settings),
    )
    if not kept:
        raise UploadIndexingError(
            f"{source} was filtered out by the document quality thresholds and cannot be indexed."
        )
    # Guarantee provenance even if the caller passed bare documents: every
    # chunk must carry the tool-ready source path for citation grounding and
    # for idempotent delete-by-source on re-upload.
    for document in kept:
        metadata = document.metadata
        metadata["source"] = source
        metadata.setdefault("kind", "upload")

    if settings.hierarchical_chunking_enabled:
        config = HierarchicalSplitConfig(
            parent_chunk_size=settings.hierarchical_parent_chunk_size,
            child_chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
        )
        children, docstore = split_documents_hierarchical(kept, config)
        return UploadChunkResult(source=source, chunks=children, parent_docstore_additions=docstore)

    if getattr(settings, "structure_aware_chunking_enabled", False):
        kept = split_markdown_sections(kept)

    chunks = split_documents(
        kept,
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
    )
    return UploadChunkResult(source=source, chunks=chunks)


def parse_and_chunk_upload(
    path: Path,
    *,
    settings: Settings,
    source: str,
    chat_model: Any | None = None,
    chunk_context_cache: Any | None = None,
) -> UploadChunkResult:
    """Parse one saved upload, chunk it, and enforce the per-file chunk cap."""

    documents = parse_upload_document(path, source=source)
    result = chunk_upload_documents(
        documents,
        settings=settings,
        source=source,
        chat_model=chat_model,
        chunk_context_cache=chunk_context_cache,
    )
    max_chunks = max(1, int(getattr(settings, "upload_indexing_max_chunks", 2000)))
    if len(result.chunks) > max_chunks:
        raise UploadIndexingError(
            f"{source} would index {len(result.chunks)} chunks (limit {max_chunks}); not indexed."
        )
    return result


__all__ = [
    "INDEXABLE_SUFFIXES",
    "MAX_UPLOAD_DOC_CHARS",
    "UploadChunkResult",
    "UploadIndexingError",
    "chunk_upload_documents",
    "parse_and_chunk_upload",
    "parse_upload_document",
]
