from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Any

from langchain_core.tools import BaseTool, StructuredTool
from pydantic import BaseModel, Field

from ._files import FileAccessError, resolve_safe_path

if TYPE_CHECKING:
    from src.config import Settings

_SUFFIXES = (".xlsx",)
# Bound the rendered grid so a large sheet cannot flood the LLM context.
_MAX_ROWS = 100
_MAX_COLS = 20


class ExcelFileInput(BaseModel):
    """Input schema for the Excel (.xlsx) reader."""

    path: str = Field(
        ...,
        min_length=1,
        description=(
            "Path to a .xlsx spreadsheet, relative to the configured file-read "
            "root directory. Legacy .xls is not supported."
        ),
    )
    sheet: str | None = Field(
        default=None,
        description="Optional worksheet name. Defaults to the active sheet.",
    )
    max_rows: int = Field(
        default=_MAX_ROWS,
        ge=1,
        le=1000,
        description="Maximum number of rows to read from the sheet.",
    )


def build_excel_tool(
    settings: Settings,
) -> BaseTool:
    """Create an Excel .xlsx reader confined to the file-read root."""

    root = Path(settings.file_read_root)
    max_bytes = settings.file_read_max_bytes

    def _run_excel(
        path: str,
        sheet: str | None = None,
        max_rows: int = _MAX_ROWS,
    ) -> str:
        return read_excel_spreadsheet(
            path,
            root=root,
            max_bytes=max_bytes,
            sheet=sheet,
            max_rows=max_rows,
        )

    return StructuredTool.from_function(
        func=_run_excel,
        name="read_excel_spreadsheet",
        description=(
            "Read tabular data from a Microsoft Excel .xlsx spreadsheet in the "
            "local document directory. Use when the user references a "
            "spreadsheet and wants its cells/rows read, summarized, or "
            "searched. Returns sheet names and the chosen sheet's rows as a "
            "markdown table. Legacy .xls files are not supported."
        ),
        args_schema=ExcelFileInput,
    )


def read_excel_spreadsheet(
    path: str,
    *,
    root: Path,
    max_bytes: int,
    sheet: str | None = None,
    max_rows: int = _MAX_ROWS,
) -> str:
    try:
        resolved = resolve_safe_path(
            path,
            root=root,
            expected_suffixes=_SUFFIXES,
            max_bytes=max_bytes,
        )
    except FileAccessError as exc:
        return f"Could not read spreadsheet: {exc}"

    try:
        import openpyxl
    except ImportError as exc:
        return (
            "openpyxl is not installed. Run `python -m pip install openpyxl` "
            f"to enable the Excel tool: {exc}"
        )

    row_limit = max(1, int(max_rows))
    sheet_key = str(sheet).strip() if sheet is not None else ""

    def _load_and_format() -> str:
        workbook = openpyxl.load_workbook(resolved, read_only=True, data_only=True)
        try:
            sheet_names = list(workbook.sheetnames)
            worksheet = _select_sheet(workbook, sheet)
            if worksheet is None:
                available = ", ".join(sheet_names)
                return (
                    f"Sheet {sheet!r} not found in {resolved.name}. "
                    f"Available sheets: {available}."
                )
            return _format_sheet(resolved.name, sheet_names, worksheet, row_limit)
        finally:
            workbook.close()

    try:
        from ._file_cache import PARSED_FILE_CACHE

        return PARSED_FILE_CACHE.get_or_compute(
            resolved,
            parser_key=f"xlsx-render-v1:sheet={sheet_key!r}:rows={row_limit}",
            loader=_load_and_format,
        )
    except Exception as exc:
        return f"Could not open spreadsheet {resolved.name!r}: {exc}"


def _select_sheet(workbook: Any, sheet: str | None) -> Any:
    if sheet is None or not str(sheet).strip():
        return workbook.active
    name = str(sheet).strip()
    if name in workbook.sheetnames:
        return workbook[name]
    return None


def _format_sheet(
    file_name: str,
    sheet_names: list[str],
    worksheet: Any,
    max_rows: int,
) -> str:
    limit = max(1, int(max_rows))
    rows: list[list[str]] = []
    for index, row in enumerate(worksheet.iter_rows(values_only=True)):
        if index >= limit:
            break
        cells = list(row[:_MAX_COLS])
        rows.append(["" if value is None else str(value) for value in cells])

    header = (
        f"Spreadsheet {file_name} | sheets: {', '.join(sheet_names)} | "
        f"reading '{worksheet.title}':"
    )
    if not rows:
        return f"{header}\n\n(the sheet is empty)"

    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    table = _markdown_table_from_rows(normalized)
    note = ""
    if worksheet.max_row and worksheet.max_row > limit:
        note = f"\n\n(showing first {limit} of {worksheet.max_row} rows)"
    return f"{header}\n\n{table}{note}"


def _markdown_table_from_rows(rows: list[list[str]]) -> str:
    width = len(rows[0])
    header_cells = [f"C{i + 1}" for i in range(width)]
    lines = [
        "| " + " | ".join(header_cells) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for row in rows:
        safe = [cell.replace("|", "\\|").replace("\n", " ") for cell in row]
        lines.append("| " + " | ".join(safe) + " |")
    return "\n".join(lines)


__all__ = [
    "ExcelFileInput",
    "build_excel_tool",
    "read_excel_spreadsheet",
]
