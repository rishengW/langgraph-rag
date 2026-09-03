"""Session-scoped PowerPoint (.pptx) inspection and editing tools.

Edits are deliberately conservative: the model must inspect a presentation
first and provide the exact current text for every changed location.  The
uploaded source is never overwritten; a validated edited copy is published
inside the current chat session.
"""

from __future__ import annotations

import logging
import os
import re
import tempfile
import zipfile
from collections.abc import Iterator
from copy import deepcopy
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal, cast
from urllib.parse import quote

from langchain_core.tools import BaseTool, StructuredTool
from pydantic import BaseModel, Field, model_validator

from ._files import FileAccessError, resolve_safe_path

if TYPE_CHECKING:
    from src.config import Settings

logger = logging.getLogger(__name__)

PPTX_SUFFIXES = (".pptx",)
PPTX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_MAX_OPERATIONS = 50
_MAX_INSPECT_CHARS = 30_000
_MAX_TEXT_CHARS = 10_000
_MAX_ZIP_ENTRIES = 1024
_MAX_UNCOMPRESSED_BYTES = 100_000_000
_MAX_COMPRESSION_RATIO = 200
_MAX_OUTPUT_STEM_CHARS = 150
_MAX_PORTABLE_PATH_CHARS = 240
_SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")
_THREAD_ID_PATTERN = re.compile(r"^[A-Za-z0-9._-]{1,64}$")


class PowerPointEditError(Exception):
    """Raised when a presentation cannot be safely inspected or edited."""


class PowerPointEditResult:
    def __init__(self, content: str, artifact: dict[str, object] | None = None) -> None:
        self.content = content
        self.artifact = artifact


class PowerPointEditOperation(BaseModel):
    action: Literal["replace_text", "append_text", "delete_text", "replace_table_cell"]
    slide_index: int = Field(..., ge=0, description="Zero-based slide number from inspect_powerpoint.")
    shape_path: list[int] = Field(
        ...,
        min_length=1,
        max_length=8,
        description=(
            "Shape path from inspect_powerpoint. Top-level shapes use [0], and "
            "children inside groups use paths such as [2, 0]."
        ),
    )
    expected_text: str | None = Field(
        default=None,
        max_length=_MAX_TEXT_CHARS,
        description="Current target text from inspect_powerpoint; whitespace is normalized for comparison.",
    )
    new_text: str | None = Field(
        default=None,
        max_length=_MAX_TEXT_CHARS,
        description="Replacement or appended text.",
    )
    row_index: int | None = Field(default=None, ge=0)
    column_index: int | None = Field(default=None, ge=0)

    @model_validator(mode="after")
    def validate_operation(self) -> PowerPointEditOperation:
        if any(index < 0 for index in self.shape_path):
            raise ValueError("shape_path indexes must be non-negative.")
        if self.expected_text is None:
            raise ValueError("expected_text is required for every PowerPoint edit.")
        if self.action in ("replace_text", "append_text", "replace_table_cell") and self.new_text is None:
            raise ValueError(f"{self.action} requires new_text.")
        if self.action == "replace_table_cell" and (self.row_index is None or self.column_index is None):
            raise ValueError("replace_table_cell requires row_index and column_index.")
        return self


class PowerPointInspectInput(BaseModel):
    path: str = Field(..., min_length=1, description="Path to a .pptx uploaded to this chat session.")
    max_chars: int = Field(default=_MAX_INSPECT_CHARS, ge=1, le=200_000)


class PowerPointEditInput(BaseModel):
    path: str = Field(..., min_length=1, description="Path to a .pptx uploaded to this chat session.")
    operations: list[PowerPointEditOperation] = Field(..., min_length=1, max_length=_MAX_OPERATIONS)
    output_name: str | None = Field(default=None, max_length=200)


def build_powerpoint_edit_tools(
    settings: Settings, *, session_root: Path | None = None, thread_id: str = ""
) -> list[BaseTool]:
    if not (settings.file_read_enabled and settings.powerpoint_edit_enabled):
        return []
    if session_root is None or not _THREAD_ID_PATTERN.fullmatch(thread_id):
        return []
    file_root = Path(settings.file_read_root)

    def _inspect(path: str, max_chars: int = _MAX_INSPECT_CHARS) -> str:
        return inspect_powerpoint(path, session_root=session_root, file_root=file_root,
                                  max_bytes=settings.file_read_max_bytes, max_chars=max_chars)

    def _edit(path: str, operations: list[PowerPointEditOperation], output_name: str | None = None) -> tuple[str, dict[str, object] | None]:
        result = edit_powerpoint(path, operations=operations, session_root=session_root,
                                 file_root=file_root, max_bytes=settings.file_read_max_bytes,
                                 thread_id=thread_id, output_name=output_name)
        return result.content, result.artifact

    return [
        StructuredTool.from_function(
            func=_inspect, name="inspect_powerpoint",
            description=("List slides, shapes, text, and table cells in an uploaded .pptx. "
                         "Call this before edit_powerpoint to obtain exact indexes and expected_text."),
            args_schema=PowerPointInspectInput,
        ),
        StructuredTool.from_function(
            func=_edit, name="edit_powerpoint",
            description=("Apply structured text or table-cell edits to an uploaded .pptx and "
                         "save a new downloadable copy. The original is never changed. "
                         "Always call inspect_powerpoint first and pass its exact expected_text."),
            args_schema=PowerPointEditInput, response_format="content_and_artifact",
        ),
    ]


def inspect_powerpoint(path: str, *, session_root: Path, file_root: Path, max_bytes: int,
                       max_chars: int = _MAX_INSPECT_CHARS) -> str:
    try:
        resolved = _resolve_session_pptx(path, session_root=session_root, file_root=file_root, max_bytes=max_bytes)
        presentation = _load_presentation(resolved)
    except PowerPointEditError as exc:
        return f"Could not inspect PowerPoint presentation: {exc}"
    lines = [f"SLIDES ({len(presentation.slides)}):"]
    for slide_index, slide in enumerate(presentation.slides):
        lines.append(f"Slide {slide_index}:")
        for shape_path, shape in _iter_shapes(slide.shapes):
            path_label = "/".join(str(index) for index in shape_path)
            indent = "  " * len(shape_path)
            text = _normalize(getattr(shape, "text", ""))
            if getattr(shape, "has_table", False):
                lines.append(
                    f"{indent}[{path_label}] TABLE "
                    f"({len(shape.table.rows)}x{len(shape.table.columns)})"
                )
                for row_index, row in enumerate(shape.table.rows):
                    for column_index, cell in enumerate(row.cells):
                        cell_text = _normalize(cell.text)
                        lines.append(
                            f"{indent}  [table row {row_index} col {column_index}] "
                            f"{cell_text or '(empty)'}"
                        )
            else:
                kind = getattr(shape, "shape_type", "shape")
                lines.append(f"{indent}[{path_label}] {kind}: {text or '(empty)'}")
    body = "\n".join(lines)
    limit = max(1, int(max_chars))
    if len(body) > limit:
        return f"Structure of {resolved.name} (truncated to {limit:,} of {len(body):,} chars):\n\n{body[:limit]}"
    return f"Structure of {resolved.name}:\n\n{body}"


def edit_powerpoint(path: str, *, operations: list[PowerPointEditOperation], session_root: Path,
                    file_root: Path, max_bytes: int, thread_id: str = "",
                    output_name: str | None = None) -> PowerPointEditResult:
    source_name = "(unresolved)"
    try:
        resolved = _resolve_session_pptx(path, session_root=session_root, file_root=file_root, max_bytes=max_bytes)
        source_name = resolved.name
        presentation = _load_presentation(resolved)
        _apply_operations(presentation, operations)
        published = _publish(presentation, session_root=session_root, source_path=resolved,
                             output_name=output_name, max_bytes=max_bytes)
    except PowerPointEditError as exc:
        logger.info("powerpoint_edit failed: thread=%s source=%s reason=%s", thread_id or "(none)", source_name, type(exc).__name__)
        return PowerPointEditResult(f"Could not edit PowerPoint presentation: {exc}")
    summary = ", ".join(op.action for op in operations)
    size = published.stat().st_size
    return PowerPointEditResult(
        f"Applied {len(operations)} edit(s) ({summary}) to {source_name}. The original file is unchanged; the edited copy was saved as {published.name} and is available to download.",
        build_file_artifact(thread_id=thread_id, filename=published.name, size_bytes=size),
    )


def _apply_operations(presentation: Any, operations: list[PowerPointEditOperation]) -> None:
    for operation in operations:
        try:
            slide = presentation.slides[operation.slide_index]
            shape = _resolve_shape(slide, operation.shape_path)
        except (IndexError, TypeError) as exc:
            raise PowerPointEditError(f"invalid slide or shape path in {operation.action}.") from exc
        actual = _shape_text(shape, operation)
        if _normalize(actual) != _normalize(operation.expected_text or ""):
            raise PowerPointEditError(
                f"expected_text does not match slide {operation.slide_index}, shape {operation.shape_path} "
                f"(found {actual!r}). Re-run inspect_powerpoint and retry."
            )
        if operation.action == "replace_table_cell":
            if not getattr(shape, "has_table", False):
                raise PowerPointEditError("replace_table_cell target is not a table.")
            try:
                _replace_text_frame_text(
                    shape.table.cell(operation.row_index, operation.column_index).text_frame,
                    operation.new_text or "",
                )
            except IndexError as exc:
                raise PowerPointEditError("table row or column index is out of range.") from exc
        elif operation.action == "replace_text":
            _replace_shape_text(shape, operation.new_text or "")
        elif operation.action == "append_text":
            if not getattr(shape, "has_text_frame", False):
                raise PowerPointEditError("append_text target has no text frame.")
            _append_text_preserving_format(shape.text_frame, operation.new_text or "")
        else:
            _replace_shape_text(shape, "")


def _shape_text(shape: Any, operation: PowerPointEditOperation) -> str:
    if operation.action == "replace_table_cell" and getattr(shape, "has_table", False):
        try:
            return cast(
                str,
                shape.table.cell(operation.row_index, operation.column_index).text,
            )
        except IndexError as exc:
            raise PowerPointEditError("table row or column index is out of range.") from exc
    return getattr(shape, "text", "") or ""


def _iter_shapes(
    shapes: Any, prefix: tuple[int, ...] = ()
) -> Iterator[tuple[tuple[int, ...], Any]]:
    """Yield top-level and grouped shapes with stable paths for tool calls."""

    for index, shape in enumerate(shapes):
        path = (*prefix, index)
        yield path, shape
        children = getattr(shape, "shapes", None)
        if children is not None:
            yield from _iter_shapes(children, path)


def _resolve_shape(slide: Any, shape_path: list[int]) -> Any:
    """Resolve a shape path through nested group shapes."""

    shapes = slide.shapes
    shape: Any = None
    for depth, index in enumerate(shape_path):
        shape = shapes[index]
        if depth < len(shape_path) - 1:
            children = getattr(shape, "shapes", None)
            if children is None:
                raise IndexError("shape path enters a non-group shape")
            shapes = children
    return shape


def _replace_shape_text(shape: Any, text: str) -> None:
    if not getattr(shape, "has_text_frame", False):
        raise PowerPointEditError("target shape has no text frame.")
    _replace_text_frame_text(shape.text_frame, text)


def _replace_text_frame_text(text_frame: Any, text: str) -> None:
    """Replace all text while retaining the first run and paragraph styling."""

    paragraphs = list(text_frame.paragraphs)
    if not paragraphs:
        text_frame.text = text
        return

    first = paragraphs[0]
    runs = list(first.runs)
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        first.add_run().text = text

    for paragraph in paragraphs[1:]:
        element = paragraph._p
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def _append_text_preserving_format(text_frame: Any, text: str) -> None:
    """Append a paragraph using the previous paragraph/run formatting when available."""

    template = text_frame.paragraphs[-1] if text_frame.paragraphs else None
    paragraph = text_frame.add_paragraph()
    if template is not None and template._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    run = paragraph.add_run()
    run.text = text
    if template is not None and template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))


def _resolve_session_pptx(raw_path: str, *, session_root: Path, file_root: Path, max_bytes: int) -> Path:
    text = (raw_path or "").strip().strip('"').strip("'")
    if not text:
        raise PowerPointEditError("a non-empty file path is required.")
    try:
        session = session_root.expanduser().resolve()
    except OSError as exc:
        raise PowerPointEditError(f"could not resolve the session directory: {exc}") from exc
    candidate = Path(text).expanduser()
    candidates = [candidate] if candidate.is_absolute() else [session / candidate, file_root / candidate]
    access_errors: list[FileAccessError] = []
    found_in_scope = False
    for option in candidates:
        try:
            resolved = option.resolve()
        except OSError:
            continue
        if not _is_within(resolved, session):
            continue
        found_in_scope = True
        try:
            return resolve_safe_path(
                str(resolved), root=session, expected_suffixes=PPTX_SUFFIXES, max_bytes=max_bytes
            )
        except FileAccessError as exc:
            access_errors.append(exc)
    if not found_in_scope:
        raise PowerPointEditError("access denied: the presentation is not in this chat session.")
    if access_errors:
        raise PowerPointEditError(str(access_errors[-1])) from access_errors[-1]
    raise PowerPointEditError(f"could not resolve {text!r} inside this chat session.")


def _load_presentation(path: Path) -> Any:
    try:
        _guard_archive(path)
        from pptx import Presentation
        return Presentation(str(path))
    except ImportError as exc:
        raise PowerPointEditError("python-pptx is not installed; install python-pptx to enable PowerPoint editing.") from exc
    except PowerPointEditError:
        raise
    except Exception as exc:
        raise PowerPointEditError(f"invalid or unreadable .pptx file: {exc}") from exc


def _guard_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > _MAX_ZIP_ENTRIES:
                raise PowerPointEditError("presentation contains too many archive entries.")
            total = 0
            for info in entries:
                total += info.file_size
                if info.file_size and not info.compress_size:
                    raise PowerPointEditError("presentation has an unsafe compression ratio.")
                if info.compress_size and info.file_size / info.compress_size > _MAX_COMPRESSION_RATIO:
                    raise PowerPointEditError("presentation has an unsafe compression ratio.")
            if total > _MAX_UNCOMPRESSED_BYTES:
                raise PowerPointEditError("presentation expands beyond the safety limit.")
            if "[Content_Types].xml" not in archive.namelist() or "ppt/presentation.xml" not in archive.namelist():
                raise PowerPointEditError("file is missing required PowerPoint parts.")
    except PowerPointEditError:
        raise
    except (OSError, EOFError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise PowerPointEditError(f"file is not a readable PowerPoint archive: {exc}") from exc


def _publish(presentation: Any, *, session_root: Path, source_path: Path, output_name: str | None, max_bytes: int) -> Path:
    directory = session_root.expanduser().resolve()
    source = source_path.resolve()
    if not directory.is_dir() or not _is_within(source, directory):
        raise PowerPointEditError("the output directory is outside this chat session.")
    requested = Path(output_name or f"{source.stem}.edited").name
    if requested.lower().endswith(".pptx"):
        requested = requested[:-len(".pptx")]
    portable_stem_limit = _MAX_PORTABLE_PATH_CHARS - len(str(directory)) - len("-100.pptx") - 1
    if portable_stem_limit < 1:
        raise PowerPointEditError("the session directory path is too long for a safe output file.")
    stem_limit = min(_MAX_OUTPUT_STEM_CHARS, portable_stem_limit)
    stem = _SAFE_NAME.sub("_", requested.strip().strip("."))[:stem_limit]
    if not stem:
        raise PowerPointEditError("a usable output filename is required.")
    target: Path | None = None
    for variant in range(1, 101):
        name = f"{stem}.pptx" if variant == 1 else f"{stem}-{variant}.pptx"
        candidate = directory / name
        if candidate.resolve() == source:
            continue
        try:
            candidate.touch(exist_ok=False)
            target = candidate
            break
        except FileExistsError:
            continue
        except OSError as exc:
            raise PowerPointEditError(f"could not reserve the output file: {exc}") from exc
    if target is None:
        raise PowerPointEditError("too many edited copies with this name exist in the session.")
    try:
        handle, temp_name = tempfile.mkstemp(
            prefix=".powerpoint_edit-", suffix=".pptx.tmp", dir=str(directory)
        )
        os.close(handle)
    except OSError as exc:
        target.unlink(missing_ok=True)
        raise PowerPointEditError(f"could not create a temporary file: {exc}") from exc
    temp = Path(temp_name)
    try:
        presentation.save(str(temp))
        if temp.stat().st_size > max_bytes:
            raise PowerPointEditError("the edited presentation exceeds the file size limit.")
        _load_presentation(temp)
        os.replace(temp, target)
    except PowerPointEditError:
        temp.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise
    except (OSError, ValueError) as exc:
        temp.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise PowerPointEditError(f"could not save the edited presentation: {exc}") from exc
    except Exception as exc:
        temp.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise PowerPointEditError(f"could not save the edited presentation: {exc}") from exc
    return target


def build_file_artifact(*, thread_id: str, filename: str, size_bytes: int) -> dict[str, object]:
    return {"type": "file", "version": 1, "kind": "download", "provider": "chat_upload",
            "threadId": thread_id, "filename": filename, "mimeType": PPTX_MIME_TYPE,
            "sizeBytes": size_bytes, "url": f"/chat/{quote(thread_id, safe='')}/files/{quote(filename, safe='')}"}


def _normalize(value: str) -> str:
    return " ".join((value or "").split())


def _is_within(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


__all__ = ["PPTX_MIME_TYPE", "PowerPointEditInput", "PowerPointEditOperation", "PowerPointInspectInput",
           "PowerPointEditResult", "build_powerpoint_edit_tools", "edit_powerpoint", "inspect_powerpoint"]
