"""Session-scoped Word (.docx) creation, inspection, and editing tools.

Writes are deliberately narrower than the read-only reader in ``word_file.py``:

* **Session-bound.** Files are confined to the current chat thread's upload
  directory (``chat_uploads/<thread-id>/``), not the global ``file_read_root``.
  An agent therefore cannot touch another session's documents.
* **No overwrite.** Creations pick a collision-safe filename. Edits write a
  new ``<name>.edited.docx`` beside the source, which is never modified.
* **Validated publication.** Every output is saved to a temporary file,
  re-opened to prove it is a valid DOCX, then published atomically.
* **Optimistic concurrency.** Every destructive operation carries the
  ``expected_text`` the model believes is at the target location. If any
  operation's expectation does not match, the whole batch is rejected and
  nothing is written.

Uploaded documents are untrusted input: the archive is checked for ZIP-bomb
characteristics before python-docx is allowed to parse it.
"""

from __future__ import annotations

import logging
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal
from urllib.parse import quote

from langchain_core.tools import BaseTool, StructuredTool
from pydantic import BaseModel, Field, model_validator

from ._files import FileAccessError, resolve_safe_path

if TYPE_CHECKING:
    from ..config import Settings

logger = logging.getLogger(__name__)

_SUFFIXES = (".docx",)
_INSPECT_MAX_CHARS = 20_000
_MAX_OPERATIONS = 50
_MAX_TEXT_CHARS = 10_000
_MAX_CREATE_BLOCKS = 200
_MAX_CREATE_ITEMS = 100
_MAX_CREATE_ROWS = 200
_MAX_CREATE_COLUMNS = 12
_MAX_CREATE_TEXT_CHARS = 200_000

# ZIP-bomb guards, applied before python-docx parses an uploaded archive.
_MAX_ZIP_ENTRIES = 512
_MAX_UNCOMPRESSED_BYTES = 50_000_000
_MAX_COMPRESSION_RATIO = 200

# Parts every real .docx contains; a zip missing them is not a Word document.
_REQUIRED_PARTS = ("[Content_Types].xml", "word/document.xml")

DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Artifact envelope constants, mirrored by the validator in graph/artifacts.py.
FILE_ARTIFACT_TYPE = "file"
FILE_ARTIFACT_VERSION = 1
FILE_ARTIFACT_KIND = "download"
FILE_ARTIFACT_PROVIDER = "chat_upload"

_SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")
_THREAD_ID_PATTERN = re.compile(r"^[A-Za-z0-9._-]{1,64}$")
_MAX_OUTPUT_VARIANTS = 100


class WordEditError(Exception):
    """Raised when a document cannot be inspected or edited safely."""


@dataclass(frozen=True)
class WordEditResult:
    """Internal content/artifact result for a document edit."""

    content: str
    artifact: dict[str, object] | None = None


class WordEditOperation(BaseModel):
    """One structured edit applied to a Word document.

    ``expected_text`` is the text the model believes currently occupies the
    target location. It is compared with whitespace collapsed, so the model
    does not have to reproduce runs of spaces byte-for-byte, but it must have
    actually read the location before changing it.
    """

    action: Literal[
        "replace_paragraph",
        "append_paragraph",
        "delete_paragraph",
        "replace_table_cell",
    ] = Field(..., description="The kind of edit to apply.")
    paragraph_index: int | None = Field(
        default=None,
        ge=0,
        description=(
            "Zero-based paragraph number from inspect_word_document. Required "
            "for replace_paragraph and delete_paragraph."
        ),
    )
    table_index: int | None = Field(
        default=None,
        ge=0,
        description="Zero-based table number. Required for replace_table_cell.",
    )
    row_index: int | None = Field(
        default=None,
        ge=0,
        description="Zero-based row number in the table. Required for replace_table_cell.",
    )
    column_index: int | None = Field(
        default=None,
        ge=0,
        description="Zero-based column number in the row. Required for replace_table_cell.",
    )
    expected_text: str | None = Field(
        default=None,
        max_length=_MAX_TEXT_CHARS,
        description=(
            "The text currently at the target location, as reported by "
            "inspect_word_document. Required for every action except "
            "append_paragraph. The edit is rejected if it does not match."
        ),
    )
    new_text: str | None = Field(
        default=None,
        max_length=_MAX_TEXT_CHARS,
        description=(
            "The replacement or new text. Required for replace_paragraph, "
            "append_paragraph, and replace_table_cell."
        ),
    )

    @model_validator(mode="after")
    def _check_required_fields(self) -> WordEditOperation:
        """Reject operations missing the fields their action needs."""

        missing: list[str] = []
        if self.action in ("replace_paragraph", "delete_paragraph"):
            if self.paragraph_index is None:
                missing.append("paragraph_index")
            if self.expected_text is None:
                missing.append("expected_text")
        if self.action == "replace_table_cell":
            for name in ("table_index", "row_index", "column_index"):
                if getattr(self, name) is None:
                    missing.append(name)
            if self.expected_text is None:
                missing.append("expected_text")
        if self.action in ("replace_paragraph", "append_paragraph", "replace_table_cell") and (
            self.new_text is None
        ):
            missing.append("new_text")

        if missing:
            fields = ", ".join(missing)
            raise ValueError(f"{self.action} requires: {fields}.")
        return self


class WordInspectInput(BaseModel):
    """Input schema for the Word document inspector."""

    path: str = Field(
        ...,
        min_length=1,
        description=(
            "Path to a .docx file uploaded to this chat session, exactly as "
            "listed in the upload note."
        ),
    )
    max_chars: int = Field(
        default=_INSPECT_MAX_CHARS,
        ge=1,
        le=200_000,
        description="Maximum characters of the numbered listing to return.",
    )


class WordContentBlock(BaseModel):
    """One structured block in a newly created Word document."""

    kind: Literal[
        "heading",
        "paragraph",
        "bullet_list",
        "numbered_list",
        "table",
        "page_break",
    ] = Field(..., description="The kind of document block to create.")
    text: str | None = Field(
        default=None,
        max_length=_MAX_TEXT_CHARS,
        description="Text for a heading or paragraph block.",
    )
    level: int = Field(
        default=1,
        ge=1,
        le=3,
        description="Heading level (1-3); used only for heading blocks.",
    )
    items: list[str] | None = Field(
        default=None,
        max_length=_MAX_CREATE_ITEMS,
        description="Items for a bullet_list or numbered_list block.",
    )
    rows: list[list[str]] | None = Field(
        default=None,
        max_length=_MAX_CREATE_ROWS,
        description=(
            "Rows for a table block. Every row must have the same number of "
            "columns; the first row is styled as a header by default."
        ),
    )
    has_header: bool = Field(
        default=True,
        description="Whether the first table row should be styled as a header.",
    )

    @model_validator(mode="after")
    def _check_block_shape(self) -> WordContentBlock:
        """Require the payload associated with each block kind."""

        if self.kind in ("heading", "paragraph"):
            if self.text is None or not self.text.strip():
                raise ValueError(f"{self.kind} requires non-empty text.")
        elif self.kind in ("bullet_list", "numbered_list"):
            if not self.items or any(not item.strip() for item in self.items):
                raise ValueError(f"{self.kind} requires non-empty items.")
        elif self.kind == "table":
            if not self.rows:
                raise ValueError("table requires at least one row.")
            column_count = len(self.rows[0])
            if column_count < 1 or column_count > _MAX_CREATE_COLUMNS:
                raise ValueError(
                    f"table rows must contain 1-{_MAX_CREATE_COLUMNS} columns."
                )
            if any(len(row) != column_count for row in self.rows):
                raise ValueError("every table row must have the same number of columns.")
        return self


class WordCreateInput(BaseModel):
    """Input schema for creating a professionally formatted Word document."""

    filename: str = Field(
        ...,
        min_length=1,
        max_length=200,
        description=(
            "Bare filename for the new Word document. A .docx suffix is added "
            "when needed; existing files are never overwritten."
        ),
    )
    title: str | None = Field(
        default=None,
        max_length=500,
        description="Optional document title shown at the top of the first page.",
    )
    subtitle: str | None = Field(
        default=None,
        max_length=1_000,
        description="Optional subtitle shown directly below the title.",
    )
    author: str | None = Field(
        default=None,
        max_length=200,
        description="Optional author stored in the document properties.",
    )
    blocks: list[WordContentBlock] = Field(
        ...,
        min_length=1,
        max_length=_MAX_CREATE_BLOCKS,
        description=(
            "Ordered document content: headings, paragraphs, real bullet or "
            "numbered lists, tables, and page breaks."
        ),
    )

    @model_validator(mode="after")
    def _check_document_text(self) -> WordCreateInput:
        """Reject unsafe XML characters and excessive aggregate content."""

        values: list[str] = [self.title or "", self.subtitle or "", self.author or ""]
        for block in self.blocks:
            values.append(block.text or "")
            values.extend(block.items or [])
            values.extend(cell for row in (block.rows or []) for cell in row)

        total_chars = sum(len(value) for value in values)
        if total_chars > _MAX_CREATE_TEXT_CHARS:
            raise ValueError(
                f"document content exceeds the {_MAX_CREATE_TEXT_CHARS:,}-character limit."
            )
        if any(_contains_invalid_xml_characters(value) for value in values):
            raise ValueError("document text contains unsupported control characters.")
        return self


class WordEditInput(BaseModel):
    """Input schema for the Word document editor."""

    path: str = Field(
        ...,
        min_length=1,
        description=(
            "Path to a .docx file uploaded to this chat session, exactly as "
            "listed in the upload note. This file is never modified."
        ),
    )
    operations: list[WordEditOperation] = Field(
        ...,
        min_length=1,
        max_length=_MAX_OPERATIONS,
        description=(
            "The edits to apply, in order. All of them are checked against the "
            "document first; if any expected_text does not match, none are applied."
        ),
    )
    output_name: str | None = Field(
        default=None,
        max_length=200,
        description=(
            "Optional file name for the edited copy. Defaults to "
            "'<original>.edited.docx'. A .docx suffix is enforced."
        ),
    )


def build_word_edit_tools(
    settings: Settings,
    *,
    session_root: Path | None = None,
    thread_id: str = "",
) -> list[BaseTool]:
    """Create the session-scoped Word inspect/edit tools.

    Returns an empty list unless both file reading and Word editing are
    enabled *and* a session upload directory was supplied. Without a session
    root there is no safe place to confine the editor, so no tool is offered.
    """

    if not (settings.file_read_enabled and settings.word_edit_enabled):
        return []
    if session_root is None or not _THREAD_ID_PATTERN.fullmatch(thread_id):
        return []

    file_root = Path(settings.file_read_root)
    max_bytes = settings.file_read_max_bytes

    def _run_create(
        filename: str,
        blocks: list[WordContentBlock],
        title: str | None = None,
        subtitle: str | None = None,
        author: str | None = None,
    ) -> tuple[str, dict[str, object] | None]:
        result = create_word_document(
            filename,
            blocks=blocks,
            session_root=session_root,
            file_root=file_root,
            max_bytes=max_bytes,
            thread_id=thread_id,
            title=title,
            subtitle=subtitle,
            author=author,
        )
        return result.content, result.artifact

    def _run_inspect(path: str, max_chars: int = _INSPECT_MAX_CHARS) -> str:
        return inspect_word_document(
            path,
            session_root=session_root,
            file_root=file_root,
            max_bytes=max_bytes,
            max_chars=max_chars,
        )

    def _run_edit(
        path: str,
        operations: list[WordEditOperation],
        output_name: str | None = None,
    ) -> tuple[str, dict[str, object] | None]:
        result = edit_word_document(
            path,
            operations=operations,
            session_root=session_root,
            file_root=file_root,
            max_bytes=max_bytes,
            thread_id=thread_id,
            output_name=output_name,
        )
        return result.content, result.artifact

    create_tool = StructuredTool.from_function(
        func=_run_create,
        name="create_word_document",
        description=(
            "Create a professionally formatted Word .docx file in the current "
            "chat session and return it as a download. Content is supplied as "
            "structured headings, paragraphs, real bullet or numbered lists, "
            "tables, and page breaks. Existing files are never overwritten; a "
            "numbered filename is chosen on collision. Use only when the user "
            "explicitly asks to create a Word document."
        ),
        args_schema=WordCreateInput,
        response_format="content_and_artifact",
    )

    inspect_tool = StructuredTool.from_function(
        func=_run_inspect,
        name="inspect_word_document",
        description=(
            "List the numbered paragraphs and table cells of a .docx file "
            "uploaded to this chat session. Use this BEFORE edit_word_document: "
            "the numbers and exact text it returns are what the edit tool needs "
            "to target a location. Returns paragraph indexes, table/row/column "
            "indexes, and the current text of each."
        ),
        args_schema=WordInspectInput,
    )

    edit_tool = StructuredTool.from_function(
        func=_run_edit,
        name="edit_word_document",
        description=(
            "Apply structured edits to a .docx file uploaded to this chat "
            "session and save the result as a NEW downloadable file; the "
            "original upload is never changed. Supports replacing a paragraph, "
            "appending a paragraph, deleting a paragraph, and replacing a table "
            "cell. Call inspect_word_document first to get the indexes and the "
            "exact current text, which you must pass as expected_text. Use only "
            "when the user explicitly asks for the document to be changed."
        ),
        args_schema=WordEditInput,
        response_format="content_and_artifact",
    )

    return [create_tool, inspect_tool, edit_tool]


def create_word_document(
    filename: str,
    *,
    blocks: list[WordContentBlock],
    session_root: Path,
    file_root: Path,
    max_bytes: int,
    thread_id: str = "",
    title: str | None = None,
    subtitle: str | None = None,
    author: str | None = None,
) -> WordEditResult:
    """Create a styled, session-scoped DOCX and downloadable artifact."""

    output_name = "(unresolved)"
    try:
        request = WordCreateInput(
            filename=filename,
            blocks=blocks,
            title=title,
            subtitle=subtitle,
            author=author,
        )
        if not _THREAD_ID_PATTERN.fullmatch(thread_id):
            raise WordEditError("a valid chat thread is required to create a document.")
        directory = _prepare_session_directory(
            session_root=session_root,
            file_root=file_root,
        )
        stem = _creation_stem(request.filename)
        document = _build_created_document(
            blocks=request.blocks,
            title=request.title,
            subtitle=request.subtitle,
            author=request.author,
        )
        target = _reserve_created_output_path(directory=directory, stem=stem)
        output_name = target.name
        published = _write_created_document(
            document,
            target=target,
            max_bytes=max_bytes,
        )
        size_bytes = published.stat().st_size
    except (OSError, ValueError, WordEditError) as exc:
        logger.info(
            "word_create failed: thread=%s output=%s reason=%s",
            thread_id or "(none)",
            output_name,
            type(exc).__name__,
        )
        return WordEditResult(content=f"Could not create Word document: {exc}")

    logger.info(
        "word_create succeeded: thread=%s output=%s blocks=%d bytes=%d",
        thread_id,
        published.name,
        len(request.blocks),
        size_bytes,
    )
    return WordEditResult(
        content=(
            f"Created {published.name} in this chat session with "
            f"{len(request.blocks)} content block(s). The Word document is available "
            "to download."
        ),
        artifact=build_file_artifact(
            thread_id=thread_id,
            filename=published.name,
            size_bytes=size_bytes,
        ),
    )


def _contains_invalid_xml_characters(value: str) -> bool:
    """Return whether text contains a character XML 1.0 cannot represent."""

    return any(
        not (
            code in (0x09, 0x0A, 0x0D)
            or 0x20 <= code <= 0xD7FF
            or 0xE000 <= code <= 0xFFFD
            or 0x10000 <= code <= 0x10FFFF
        )
        for code in map(ord, value)
    )


def _prepare_session_directory(*, session_root: Path, file_root: Path) -> Path:
    """Create and validate the current session directory under the file root."""

    try:
        root = file_root.expanduser().resolve()
        directory = session_root.expanduser().resolve()
    except OSError as exc:
        raise WordEditError(f"could not resolve the session directory: {exc}") from exc
    if directory == root or not _is_within(directory, root):
        raise WordEditError("the output directory is outside the configured file root.")
    try:
        directory.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise WordEditError(f"could not create the session directory: {exc}") from exc
    if not directory.is_dir():
        raise WordEditError("the session output path is not a directory.")
    return directory


def _creation_stem(filename: str) -> str:
    """Sanitize a requested creation name and return its filename stem."""

    requested = Path(str(filename or "")).name
    if requested.lower().endswith(".docx"):
        requested = requested[: -len(".docx")]
    cleaned = _SAFE_NAME.sub("_", requested.strip().strip("."))[:150]
    if not cleaned:
        raise WordEditError("a usable .docx filename is required.")
    return cleaned


def _reserve_created_output_path(*, directory: Path, stem: str) -> Path:
    """Exclusively reserve a collision-safe filename for a new document."""

    for variant in range(1, _MAX_OUTPUT_VARIANTS + 1):
        name = f"{stem}.docx" if variant == 1 else f"{stem}-{variant}.docx"
        candidate = directory / name
        try:
            candidate.touch(exist_ok=False)
            return candidate
        except FileExistsError:
            continue
        except OSError as exc:
            raise WordEditError(f"could not create the output file: {exc}") from exc
    raise WordEditError(
        f"too many files named {stem!r} already exist in this session; "
        "download or remove some before creating another."
    )


def _build_created_document(
    *,
    blocks: list[WordContentBlock],
    title: str | None,
    subtitle: str | None,
    author: str | None,
) -> Any:
    """Build a DOCX using the explicit standard-business style preset."""

    try:
        import docx
        from docx.enum.section import WD_ORIENT
        from docx.shared import Inches
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise WordEditError(
            "python-docx is not installed; run 'pip install python-docx' to "
            "enable Word document creation."
        ) from exc

    document = docx.Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    _configure_created_styles(document)
    _configure_page_footer(section)
    document.core_properties.title = (title or _title_from_blocks(blocks)).strip()
    if author and author.strip():
        document.core_properties.author = author.strip()

    if title and title.strip():
        paragraph = document.add_paragraph(style="Title")
        paragraph.add_run(title.strip())
    if subtitle and subtitle.strip():
        paragraph = document.add_paragraph(style="Subtitle")
        paragraph.add_run(subtitle.strip())

    for block in blocks:
        if block.kind == "heading":
            paragraph = document.add_paragraph(style=f"Heading {block.level}")
            paragraph.add_run((block.text or "").strip())
        elif block.kind == "paragraph":
            document.add_paragraph((block.text or "").strip(), style="Normal")
        elif block.kind in ("bullet_list", "numbered_list"):
            num_id = _add_numbering_definition(
                document,
                ordered=block.kind == "numbered_list",
            )
            for item in block.items or []:
                paragraph = document.add_paragraph(style="Normal")
                paragraph.add_run(item.strip())
                _apply_list_numbering(paragraph, num_id=num_id)
        elif block.kind == "table":
            _add_created_table(
                document,
                rows=block.rows or [],
                has_header=block.has_header,
            )
        else:
            _add_created_page_break(document)

    return document


def _title_from_blocks(blocks: list[WordContentBlock]) -> str:
    """Return a useful core-property title when no visible title is supplied."""

    for block in blocks:
        if block.kind == "heading" and block.text:
            return block.text
    return "Word document"


def _configure_created_styles(document: Any) -> None:
    """Apply exact typography and paragraph rhythm to named Word styles."""

    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    styles = document.styles

    def configure(
        name: str,
        *,
        font: str = "Calibri",
        size: float,
        color: str,
        bold: bool = False,
        italic: bool = False,
        before: float = 0,
        after: float = 0,
        line_spacing: float = 1.0,
        keep_with_next: bool = False,
        alignment: Any = WD_ALIGN_PARAGRAPH.LEFT,
    ) -> None:
        style = styles[name]
        style.font.name = font
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold
        style.font.italic = italic
        paragraph = style.paragraph_format
        paragraph.alignment = alignment
        paragraph.space_before = Pt(before)
        paragraph.space_after = Pt(after)
        paragraph.line_spacing = line_spacing
        paragraph.keep_with_next = keep_with_next
        paragraph.widow_control = True

    configure("Normal", size=11, color="202124", after=6, line_spacing=1.10)
    configure(
        "Title",
        size=24,
        color="1F3A5F",
        bold=True,
        after=4,
        line_spacing=1.0,
        keep_with_next=True,
    )
    configure(
        "Subtitle",
        size=12,
        color="5F6368",
        italic=True,
        after=18,
        line_spacing=1.0,
        keep_with_next=True,
    )
    configure(
        "Heading 1",
        size=16,
        color="2E74B5",
        bold=True,
        before=16,
        after=8,
        line_spacing=1.0,
        keep_with_next=True,
    )
    configure(
        "Heading 2",
        size=13,
        color="2E74B5",
        bold=True,
        before=12,
        after=6,
        line_spacing=1.0,
        keep_with_next=True,
    )
    configure(
        "Heading 3",
        size=12,
        color="1F4D78",
        bold=True,
        before=8,
        after=4,
        line_spacing=1.0,
        keep_with_next=True,
    )


def _configure_page_footer(section: Any) -> None:
    """Add a restrained, right-aligned PAGE field to the section footer."""

    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    prefix = paragraph.add_run("Page ")
    prefix.font.name = "Calibri"
    prefix.font.size = Pt(9)
    prefix.font.color.rgb = RGBColor.from_string("6B7280")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "6B7280")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.extend((color, size))
    text = OxmlElement("w:t")
    text.text = "1"
    run.extend((properties, text))
    field.append(run)
    paragraph._p.append(field)


def _add_created_page_break(document: Any) -> None:
    """Append a Word-native page break without relying on untyped convenience APIs."""

    from docx.enum.text import WD_BREAK

    paragraph = document.add_paragraph(style="Normal")
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _add_numbering_definition(document: Any, *, ordered: bool) -> int:
    """Create a real level-zero list definition and return its numId."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "decimal" if ordered else "bullet")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if ordered else "\u2022")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    paragraph_properties.extend((tabs, indent, spacing))
    level.extend((start, num_format, level_text, justification, paragraph_properties))
    abstract.append(level)
    numbering.append(abstract)

    instance = OxmlElement("w:num")
    instance.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    instance.append(reference)
    numbering.append(instance)
    return num_id


def _apply_list_numbering(paragraph: Any, *, num_id: int) -> None:
    """Attach a paragraph to a real numbering definition."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    existing = properties.find(qn("w:numPr"))
    if existing is not None:
        properties.remove(existing)
    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_properties.extend((level, number))
    properties.append(num_properties)


def _add_created_table(document: Any, *, rows: list[list[str]], has_header: bool) -> None:
    """Add an explicitly sized table with fixed geometry and readable cells."""

    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    table.style = "Table Grid"
    widths = _table_column_widths(rows)
    _set_table_geometry(table, widths=widths)

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if has_header and row_index == 0:
            row_properties = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_properties.append(repeat)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string("202124")
            if has_header and row_index == 0:
                run.bold = True
                shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                if shading is None:
                    shading = OxmlElement("w:shd")
                    cell._tc.get_or_add_tcPr().append(shading)
                shading.set(qn("w:fill"), "F2F4F7")

    after = document.add_paragraph(style="Normal")
    after.paragraph_format.space_after = Pt(0)


def _table_column_widths(rows: list[list[str]]) -> list[int]:
    """Allocate the 9360-DXA content width according to column content."""

    column_count = len(rows[0])
    if column_count == 1:
        return [9360]
    weights = []
    for column in range(column_count):
        longest = max(len(str(row[column])) for row in rows)
        weights.append(max(6.0, min(40.0, longest**0.5 * 5.0)))
    total_weight = sum(weights)
    minimum = min(900, 9360 // column_count)
    flexible = 9360 - minimum * column_count
    widths = [minimum + round(flexible * weight / total_weight) for weight in weights]
    widths[-1] += 9360 - sum(widths)
    return widths


def _set_table_geometry(table: Any, *, widths: list[int]) -> None:
    """Make tblW, tblInd, tblGrid, and every tcW agree in DXA units."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = table._tbl.tblPr

    def replace_property(tag: str, attributes: dict[str, str]) -> None:
        existing = properties.find(qn(tag))
        if existing is not None:
            properties.remove(existing)
        element = OxmlElement(tag)
        for name, value in attributes.items():
            element.set(qn(name), value)
        properties.append(element)

    replace_property("w:tblW", {"w:w": "9360", "w:type": "dxa"})
    replace_property("w:tblInd", {"w:w": "120", "w:type": "dxa"})
    replace_property("w:tblLayout", {"w:type": "fixed"})

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            margins = cell_properties.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell_properties.append(margins)
            for side, amount in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
                margin = margins.find(qn(f"w:{side}"))
                if margin is None:
                    margin = OxmlElement(f"w:{side}")
                    margins.append(margin)
                margin.set(qn("w:w"), str(amount))
                margin.set(qn("w:type"), "dxa")


def _write_created_document(document: Any, *, target: Path, max_bytes: int) -> Path:
    """Save a new DOCX through a validated temporary file and publish it."""

    try:
        handle, temp_name = tempfile.mkstemp(
            prefix=".word_create-",
            suffix=".docx.tmp",
            dir=str(target.parent),
        )
    except OSError as exc:
        target.unlink(missing_ok=True)
        raise WordEditError(f"could not create a temporary file: {exc}") from exc

    os.close(handle)
    temp_path = Path(temp_name)
    try:
        document.save(str(temp_path))
        output_size = temp_path.stat().st_size
        if output_size > max_bytes:
            raise WordEditError(
                f"the new document is too large ({output_size:,} bytes; "
                f"limit {max_bytes:,} bytes)."
            )
        _guard_archive(temp_path)
        _load_document(temp_path)
        os.replace(temp_path, target)
    except WordEditError:
        temp_path.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise
    except (OSError, ValueError) as exc:
        temp_path.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise WordEditError(f"could not save the new document: {exc}") from exc
    return target


def inspect_word_document(
    path: str,
    *,
    session_root: Path,
    file_root: Path,
    max_bytes: int,
    max_chars: int = _INSPECT_MAX_CHARS,
) -> str:
    """Return a numbered listing of a document's paragraphs and table cells."""

    try:
        resolved = _resolve_session_docx(
            path,
            session_root=session_root,
            file_root=file_root,
            max_bytes=max_bytes,
        )
        document = _load_document(resolved)
    except WordEditError as exc:
        return f"Could not inspect Word document: {exc}"

    lines: list[str] = []
    paragraphs = list(document.paragraphs)
    lines.append(f"PARAGRAPHS ({len(paragraphs)}):")
    for index, paragraph in enumerate(paragraphs):
        text = _normalize(paragraph.text)
        lines.append(f"[{index}] {text}" if text else f"[{index}] (empty)")

    tables = list(document.tables)
    lines.append("")
    lines.append(f"TABLES ({len(tables)}):")
    for table_index, table in enumerate(tables):
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                text = _normalize(cell.text)
                location = f"[table {table_index} row {row_index} col {column_index}]"
                lines.append(f"{location} {text}" if text else f"{location} (empty)")

    body = "\n".join(lines)
    limit = max(1, int(max_chars))
    if len(body) > limit:
        return (
            f"Structure of {resolved.name} (truncated to {limit:,} of "
            f"{len(body):,} chars):\n\n{body[:limit]}"
        )
    return f"Structure of {resolved.name}:\n\n{body}"


def edit_word_document(
    path: str,
    *,
    operations: list[WordEditOperation],
    session_root: Path,
    file_root: Path,
    max_bytes: int,
    thread_id: str = "",
    output_name: str | None = None,
) -> WordEditResult:
    """Apply ``operations`` to a document and publish the result as a new file.

    The source document is opened read-only and edits are applied to an
    in-memory copy. Nothing is written unless every operation validates.
    """

    source_name = "(unresolved)"
    try:
        resolved = _resolve_session_docx(
            path,
            session_root=session_root,
            file_root=file_root,
            max_bytes=max_bytes,
        )
        source_name = resolved.name
        document = _load_document(resolved)
        plan = _plan_operations(document, operations)
        _apply_operations(plan)
        published = _publish_document(
            document,
            session_root=session_root,
            source_path=resolved,
            output_name=output_name,
            max_bytes=max_bytes,
        )
    except WordEditError as exc:
        # Log metadata only — document contents never reach the log.
        logger.info(
            "word_edit failed: thread=%s source=%s operations=%d reason=%s",
            thread_id or "(none)",
            source_name,
            len(operations),
            type(exc).__name__,
        )
        return WordEditResult(content=f"Could not edit Word document: {exc}")

    size_bytes = published.stat().st_size
    logger.info(
        "word_edit succeeded: thread=%s source=%s output=%s operations=%d bytes=%d",
        thread_id or "(none)",
        source_name,
        published.name,
        len(operations),
        size_bytes,
    )

    summary = ", ".join(f"{op.action}" for op in operations)
    content = (
        f"Applied {len(operations)} edit(s) ({summary}) to {resolved.name}. "
        f"The original file is unchanged; the edited copy was saved as "
        f"{published.name} and is available to download."
    )
    return WordEditResult(
        content=content,
        artifact=build_file_artifact(
            thread_id=thread_id,
            filename=published.name,
            size_bytes=size_bytes,
        ),
    )


def build_file_artifact(
    *,
    thread_id: str,
    filename: str,
    size_bytes: int,
) -> dict[str, object]:
    """Build the 'file created' artifact envelope for a published document."""

    return {
        "type": FILE_ARTIFACT_TYPE,
        "version": FILE_ARTIFACT_VERSION,
        "kind": FILE_ARTIFACT_KIND,
        "provider": FILE_ARTIFACT_PROVIDER,
        "threadId": thread_id,
        "filename": filename,
        "mimeType": DOCX_MIME_TYPE,
        "sizeBytes": size_bytes,
        "url": session_file_url(thread_id, filename),
    }


def session_file_url(thread_id: str, filename: str) -> str:
    """Return the download path for a file in a chat session's upload dir."""

    return f"/chat/{quote(thread_id, safe='')}/files/{quote(filename, safe='')}"


def _resolve_session_docx(
    raw_path: str,
    *,
    session_root: Path,
    file_root: Path,
    max_bytes: int,
) -> Path:
    """Resolve a model-supplied path to a .docx inside this session's directory.

    The upload note gives the model paths relative to ``file_read_root``
    (``chat_uploads/<thread-id>/report.docx``), but a model may also use a bare
    file name. Both are accepted; anything that lands outside the session
    directory is refused.
    """

    text = (raw_path or "").strip().strip('"').strip("'")
    if not text:
        raise WordEditError("a non-empty file path is required.")
    if text.lower().endswith((".doc", ".docm", ".dotm")):
        raise WordEditError(
            "only .docx files can be edited; legacy .doc and macro-enabled "
            "Word formats are not supported."
        )

    try:
        session_resolved = session_root.expanduser().resolve()
    except OSError as exc:
        raise WordEditError(f"could not resolve the session directory: {exc}") from exc

    candidate = Path(text).expanduser()
    if candidate.is_absolute():
        candidates = [candidate]
    else:
        # Prefer the session directory, then the path as written in the upload
        # note (relative to the file-read root).
        candidates = [session_resolved / candidate, file_root.expanduser() / candidate]

    access_errors: list[FileAccessError] = []
    found_in_scope = False
    for option in candidates:
        try:
            resolved = option.resolve()
        except OSError:
            continue
        if not _is_within(resolved, session_resolved):
            continue
        found_in_scope = True
        try:
            return resolve_safe_path(
                str(resolved),
                root=session_resolved,
                expected_suffixes=_SUFFIXES,
                max_bytes=max_bytes,
            )
        except FileAccessError as exc:
            access_errors.append(exc)

    if not found_in_scope:
        raise WordEditError(
            f"access denied: {text!r} is outside this chat session's document "
            f"directory. Only files uploaded to this session can be edited."
        )
    if access_errors:
        raise WordEditError(str(access_errors[-1])) from access_errors[-1]
    raise WordEditError(f"could not resolve {text!r} inside this chat session.")


def _is_within(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def _guard_archive(path: Path) -> None:
    """Reject archives with ZIP-bomb characteristics or a non-DOCX shape."""

    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            names = {info.filename for info in infos}
    except zipfile.BadZipFile as exc:
        raise WordEditError(f"{path.name!r} is not a readable .docx archive: {exc}") from exc
    except OSError as exc:
        raise WordEditError(f"could not open {path.name!r}: {exc}") from exc

    if len(infos) > _MAX_ZIP_ENTRIES:
        raise WordEditError(
            f"{path.name!r} contains too many parts ({len(infos)}; "
            f"limit {_MAX_ZIP_ENTRIES})."
        )

    total_uncompressed = sum(max(0, info.file_size) for info in infos)
    if total_uncompressed > _MAX_UNCOMPRESSED_BYTES:
        raise WordEditError(
            f"{path.name!r} expands to too much data "
            f"({total_uncompressed:,} bytes; limit {_MAX_UNCOMPRESSED_BYTES:,})."
        )

    for info in infos:
        if info.compress_size > 0 and info.file_size / info.compress_size > _MAX_COMPRESSION_RATIO:
            raise WordEditError(
                f"{path.name!r} contains a suspiciously compressed part "
                f"({info.filename!r}); refusing to open it."
            )

    missing = [part for part in _REQUIRED_PARTS if part not in names]
    if missing:
        raise WordEditError(
            f"{path.name!r} is not a valid Word document (missing {missing[0]!r})."
        )


def _load_document(path: Path) -> Any:
    """Open a validated .docx with python-docx."""

    _guard_archive(path)
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise WordEditError(
            "python-docx is not installed; run 'pip install python-docx' to "
            "enable Word document editing."
        ) from exc

    try:
        return docx.Document(str(path))
    except Exception as exc:
        raise WordEditError(f"could not parse {path.name!r}: {exc}") from exc


@dataclass(frozen=True)
class _PlannedEdit:
    """An operation resolved to the object it will act on.

    Resolving targets to objects before anything is applied means a delete does
    not shift the indexes that later operations in the same batch refer to.
    """

    operation: WordEditOperation
    paragraph: Any = None
    cell: Any = None
    document: Any = None


def _plan_operations(
    document: Any,
    operations: list[WordEditOperation],
) -> list[_PlannedEdit]:
    """Resolve and verify every operation against the unmodified document.

    Raises ``WordEditError`` listing every problem found, so a mismatched batch
    tells the model exactly what to re-inspect instead of failing one at a time.
    """

    paragraphs = list(document.paragraphs)
    tables = list(document.tables)
    planned: list[_PlannedEdit] = []
    problems: list[str] = []
    seen_targets: set[tuple[object, ...]] = set()

    for position, operation in enumerate(operations):
        label = f"operation {position} ({operation.action})"

        if operation.action == "append_paragraph":
            planned.append(_PlannedEdit(operation=operation, document=document))
            continue

        if operation.action in ("replace_paragraph", "delete_paragraph"):
            index = operation.paragraph_index
            assert index is not None  # guaranteed by WordEditOperation validation
            paragraph_target = ("paragraph", index)
            if paragraph_target in seen_targets:
                problems.append(f"{label}: paragraph {index} is targeted more than once.")
                continue
            seen_targets.add(paragraph_target)
            if index >= len(paragraphs):
                problems.append(
                    f"{label}: paragraph_index {index} is out of range; the "
                    f"document has {len(paragraphs)} paragraphs."
                )
                continue
            paragraph = paragraphs[index]
            actual = _normalize(paragraph.text)
            expected = _normalize(operation.expected_text or "")
            if actual != expected:
                problems.append(
                    f"{label}: expected_text does not match paragraph {index}. "
                    f"The document has changed or the index is wrong."
                )
                continue
            planned.append(_PlannedEdit(operation=operation, paragraph=paragraph))
            continue

        # replace_table_cell
        table_index = operation.table_index
        row_index = operation.row_index
        column_index = operation.column_index
        assert table_index is not None  # guaranteed by WordEditOperation validation
        assert row_index is not None
        assert column_index is not None

        cell_target = ("table_cell", table_index, row_index, column_index)
        if cell_target in seen_targets:
            problems.append(
                f"{label}: table {table_index} row {row_index} col "
                f"{column_index} is targeted more than once."
            )
            continue
        seen_targets.add(cell_target)

        if table_index >= len(tables):
            problems.append(
                f"{label}: table_index {table_index} is out of range; the "
                f"document has {len(tables)} tables."
            )
            continue
        rows = list(tables[table_index].rows)
        if row_index >= len(rows):
            problems.append(
                f"{label}: row_index {row_index} is out of range; table "
                f"{table_index} has {len(rows)} rows."
            )
            continue
        cells = list(rows[row_index].cells)
        if column_index >= len(cells):
            problems.append(
                f"{label}: column_index {column_index} is out of range; row "
                f"{row_index} has {len(cells)} cells."
            )
            continue

        cell = cells[column_index]
        actual = _normalize(cell.text)
        expected = _normalize(operation.expected_text or "")
        if actual != expected:
            problems.append(
                f"{label}: expected_text does not match table {table_index} "
                f"row {row_index} col {column_index}."
            )
            continue
        planned.append(_PlannedEdit(operation=operation, cell=cell))

    if problems:
        detail = " ".join(problems)
        raise WordEditError(
            f"no changes were made because {len(problems)} operation(s) did not "
            f"match the document: {detail} Re-run inspect_word_document and retry."
        )
    return planned


def _apply_operations(planned: list[_PlannedEdit]) -> None:
    """Apply already-validated operations to the in-memory document."""

    for item in planned:
        operation = item.operation
        if operation.action == "append_paragraph":
            item.document.add_paragraph(operation.new_text or "")
        elif operation.action == "replace_paragraph":
            _set_paragraph_text(item.paragraph, operation.new_text or "")
        elif operation.action == "delete_paragraph":
            _remove_paragraph(item.paragraph)
        else:  # replace_table_cell
            _set_cell_text(item.cell, operation.new_text or "")


def _set_paragraph_text(paragraph: Any, text: str) -> None:
    """Replace a paragraph's text, keeping the first run's formatting."""

    runs = list(paragraph.runs)
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _set_cell_text(cell: Any, text: str) -> None:
    """Replace a table cell's text, keeping the first paragraph's formatting."""

    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        cell.add_paragraph(text)
        return
    _set_paragraph_text(paragraphs[0], text)
    for extra in paragraphs[1:]:
        _remove_paragraph(extra)


def _remove_paragraph(paragraph: Any) -> None:
    """Detach a paragraph from its parent element."""

    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _publish_document(
    document: Any,
    *,
    session_root: Path,
    source_path: Path,
    output_name: str | None,
    max_bytes: int,
) -> Path:
    """Save, validate, then atomically publish the edited document.

    The output name is reserved with an exclusive create so two concurrent
    edits cannot pick the same file, and the temporary file is validated as a
    re-openable DOCX before it replaces the reservation.
    """

    target = _reserve_output_path(
        session_root=session_root,
        source_path=source_path,
        output_name=output_name,
    )

    try:
        handle, temp_name = tempfile.mkstemp(
            prefix=".word_edit-",
            suffix=".docx.tmp",
            dir=str(target.parent),
        )
    except OSError as exc:
        target.unlink(missing_ok=True)
        raise WordEditError(f"could not create a temporary file: {exc}") from exc

    os.close(handle)
    temp_path = Path(temp_name)
    try:
        document.save(str(temp_path))
        output_size = temp_path.stat().st_size
        if output_size > max_bytes:
            raise WordEditError(
                f"the edited document is too large ({output_size:,} bytes; "
                f"limit {max_bytes:,} bytes)."
            )
        _guard_archive(temp_path)
        _load_document(temp_path)
        os.replace(temp_path, target)
    except WordEditError:
        temp_path.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise
    except (OSError, ValueError) as exc:
        temp_path.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise WordEditError(f"could not save the edited document: {exc}") from exc

    return target


def _reserve_output_path(
    *,
    session_root: Path,
    source_path: Path,
    output_name: str | None,
) -> Path:
    """Exclusively create, and return, a free output path in the session dir."""

    stem = _output_stem(source_path=source_path, output_name=output_name)
    try:
        directory = session_root.expanduser().resolve()
        resolved_source = source_path.resolve()
    except OSError as exc:
        raise WordEditError(f"could not resolve the output directory: {exc}") from exc
    if not directory.is_dir() or not _is_within(resolved_source, directory):
        raise WordEditError("the output directory is outside this chat session.")

    for variant in range(1, _MAX_OUTPUT_VARIANTS + 1):
        name = f"{stem}.docx" if variant == 1 else f"{stem}-{variant}.docx"
        candidate = directory / name
        if candidate.resolve() == resolved_source:
            # Never reserve the source itself, whatever name was requested.
            continue
        try:
            candidate.touch(exist_ok=False)
            return candidate
        except FileExistsError:
            continue
        except OSError as exc:
            raise WordEditError(f"could not create the output file: {exc}") from exc

    raise WordEditError(
        f"too many edited copies of {source_path.name!r} already exist in this "
        f"session; download or remove some before editing again."
    )


def _output_stem(*, source_path: Path, output_name: str | None) -> str:
    """Return the sanitized base name (no suffix) for the edited copy."""

    if output_name:
        requested = Path(str(output_name)).name
        if requested.lower().endswith(".docx"):
            requested = requested[: -len(".docx")]
        cleaned = _SAFE_NAME.sub("_", requested.strip().strip("."))[:150]
        if cleaned:
            return cleaned
    return f"{source_path.stem}.edited"


def _normalize(text: str) -> str:
    """Collapse whitespace so text comparisons ignore formatting noise."""

    return " ".join(str(text or "").split())


__all__ = [
    "DOCX_MIME_TYPE",
    "FILE_ARTIFACT_KIND",
    "FILE_ARTIFACT_PROVIDER",
    "FILE_ARTIFACT_TYPE",
    "FILE_ARTIFACT_VERSION",
    "WordEditError",
    "WordContentBlock",
    "WordCreateInput",
    "WordEditInput",
    "WordEditOperation",
    "WordEditResult",
    "WordInspectInput",
    "build_file_artifact",
    "build_word_edit_tools",
    "create_word_document",
    "edit_word_document",
    "inspect_word_document",
    "session_file_url",
]
