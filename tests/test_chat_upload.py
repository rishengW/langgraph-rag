from __future__ import annotations

import io
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage

from src.chat import api as chat_api
from src.chat.uploads import build_upload_context_note
from src.graph.artifacts import normalize_file_artifact
from src.tools import build_text_file_tool, build_word_tool
from src.tools.excel_edit import ExcelEditOperation, XLSX_MIME_TYPE, edit_excel
from src.tools.text_edit import TextEditOperation, create_text_file, edit_text_file
from src.tools.word_edit import WordEditOperation, edit_word_document

_DOCX_DOCUMENT_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/'
    'wordprocessingml/2006/main"><w:body>'
    "<w:p><w:r><w:t>Quarterly revenue rose 12 percent.</w:t></w:r></w:p>"
    "</w:body></w:document>"
)


def _docx_bytes() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", _DOCX_DOCUMENT_XML)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    pptx = pytest.importorskip("pptx")
    buffer = io.BytesIO()
    presentation = pptx.Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(buffer)
    return buffer.getvalue()


class _FakeGraph:
    def __init__(self) -> None:
        self.messages: list = []

    def get_state(self, config):
        return SimpleNamespace(values={"messages": list(self.messages)})

    def invoke(self, inputs, config):
        self.messages.extend(inputs["messages"])
        self.messages.append(AIMessage(content="ok"))
        return {"messages": list(self.messages)}


def _client(
    monkeypatch,
    isolated_settings,
    tmp_path: Path,
    *,
    word_edit_enabled: bool = False,
    text_edit_enabled: bool = False,
    powerpoint_edit_enabled: bool = False,
    excel_edit_enabled: bool = False,
):
    settings = isolated_settings(
        source_urls=["https://chat-default.test"],
        web_search_enabled=False,
        file_read_enabled=True,
        word_edit_enabled=word_edit_enabled,
        text_edit_enabled=text_edit_enabled,
        powerpoint_edit_enabled=powerpoint_edit_enabled,
        excel_edit_enabled=excel_edit_enabled,
        file_read_root=str(tmp_path / "files"),
        file_read_max_bytes=1_000_000,
    )
    monkeypatch.setattr(chat_api, "load_settings", lambda: settings)
    monkeypatch.setattr(
        chat_api,
        "build_chat_graph",
        lambda *args, **kwargs: _FakeGraph(),
    )
    return TestClient(chat_api.create_app()), settings


def _start_thread(client: TestClient) -> str:
    start = client.post("/chat", json={"web_search": False})
    assert start.status_code == 200
    return start.json()["thread_id"]


def test_upload_context_advertises_text_editing_only_when_enabled():
    path = "chat_uploads/thread-a/notes.txt"

    read_only = build_upload_context_note([path])
    editable = build_upload_context_note([path], text_edit_enabled=True)

    assert "inspect_text_file" not in read_only
    assert "edit_text_file" not in read_only
    assert "inspect_word_document" not in editable
    assert "create_text_file" in editable
    assert "inspect_text_file" in editable
    assert "edit_text_file" in editable
    assert "explicitly asks" in editable
    assert "expected_text" in editable
    assert "never overwrite an existing file" in editable


def test_upload_context_advertises_powerpoint_editing_only_when_enabled():
    path = "chat_uploads/thread-a/report.pptx"

    read_only = build_upload_context_note([path])
    editable = build_upload_context_note([path], powerpoint_edit_enabled=True)

    assert "inspect_powerpoint" not in read_only
    assert "edit_powerpoint" not in read_only
    assert "inspect_powerpoint" in editable
    assert "edit_powerpoint" in editable
    assert "expected_text" in editable


def test_upload_context_advertises_excel_editing_only_when_enabled():
    path = "chat_uploads/thread-a/report.xlsx"

    read_only = build_upload_context_note([path])
    editable = build_upload_context_note([path], excel_edit_enabled=True)

    assert "edit_excel_spreadsheet" not in read_only
    assert "inspect_excel_spreadsheet" in editable
    assert "edit_excel_spreadsheet" in editable
    assert "expected_value" in editable


def test_upload_injects_context_into_next_turn(monkeypatch, isolated_settings, tmp_path):
    client, _ = _client(monkeypatch, isolated_settings, tmp_path)
    with client:
        thread_id = _start_thread(client)
        client.post(
            f"/chat/{thread_id}/upload",
            files={"files": ("notes.txt", b"hello from upload", "text/plain")},
        )

        # The next turn should carry a SystemMessage announcing the upload path
        # so the LLM knows the exact path to pass to the file tools.
        registry = client.app.state.session_registry
        session = registry.get(thread_id)
        graph = session.graph

        message = client.post(
            f"/chat/{thread_id}/message", json={"message": "summarize notes.txt"}
        )
        assert message.status_code == 200

        system_texts = [
            getattr(m, "content", "")
            for m in graph.messages
            if m.__class__.__name__ == "SystemMessage"
        ]
        assert any(
            f"chat_uploads/{thread_id}/notes.txt" in text for text in system_texts
        )

        # The injected system note must not leak into the user-facing history.
        history = client.get(f"/chat/{thread_id}/history")
        roles = [turn["role"] for turn in history.json()["turns"]]
        assert "system" not in roles


def test_upload_context_not_repeated_on_later_turns(monkeypatch, isolated_settings, tmp_path):
    client, _ = _client(monkeypatch, isolated_settings, tmp_path)
    with client:
        thread_id = _start_thread(client)
        client.post(
            f"/chat/{thread_id}/upload",
            files={"files": ("notes.txt", b"hi", "text/plain")},
        )
        registry = client.app.state.session_registry
        graph = registry.get(thread_id).graph

        client.post(f"/chat/{thread_id}/message", json={"message": "first"})
        client.post(f"/chat/{thread_id}/message", json={"message": "second"})

        system_count = sum(
            1 for m in graph.messages if m.__class__.__name__ == "SystemMessage"
        )
        # Announced once on the first turn after upload, not again on the second.
        assert system_count == 1


def test_upload_text_file_then_tool_reads_it(monkeypatch, isolated_settings, tmp_path):
    client, settings = _client(monkeypatch, isolated_settings, tmp_path)
    with client:
        thread_id = _start_thread(client)

        response = client.post(
            f"/chat/{thread_id}/upload",
            files={"files": ("notes.txt", b"hello from upload", "text/plain")},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["errors"] == []
        assert len(body["files"]) == 1
        saved = body["files"][0]
        assert saved["filename"] == "notes.txt"
        assert thread_id in saved["relative_path"]

        # The saved file must be readable by the file tool using the returned
        # relative path, proving the upload lands inside the file-read root.
        tool = build_text_file_tool(settings)
        result = tool.invoke({"path": saved["relative_path"]})
        assert "hello from upload" in result


def test_upload_docx_then_word_tool_extracts_text(monkeypatch, isolated_settings, tmp_path):
    client, settings = _client(monkeypatch, isolated_settings, tmp_path)
    with client:
        thread_id = _start_thread(client)

        response = client.post(
            f"/chat/{thread_id}/upload",
            files={
                "files": (
                    "report.docx",
                    _docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        assert response.status_code == 200
        saved = response.json()["files"][0]

        tool = build_word_tool(settings)
        result = tool.invoke({"path": saved["relative_path"]})
        assert "Quarterly revenue rose 12 percent." in result


def test_upload_rejects_disallowed_type(monkeypatch, isolated_settings, tmp_path):
    client, _ = _client(monkeypatch, isolated_settings, tmp_path)
    with client:
        thread_id = _start_thread(client)

        response = client.post(
            f"/chat/{thread_id}/upload",
            files={"files": ("evil.exe", b"MZ...", "application/octet-stream")},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["files"] == []
        assert len(body["errors"]) == 1
        assert "unsupported file type" in body["errors"][0].lower()


def test_upload_powerpoint_requires_editor_flag(monkeypatch, isolated_settings, tmp_path):
    client, _ = _client(monkeypatch, isolated_settings, tmp_path)
    with client:
        thread_id = _start_thread(client)
        response = client.post(
            f"/chat/{thread_id}/upload",
            files={"files": ("report.pptx", _pptx_bytes(), "application/octet-stream")},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["files"] == []
        assert len(body["errors"]) == 1
        assert "powerpoint uploads are disabled" in body["errors"][0].lower()


def test_upload_powerpoint_when_editor_is_enabled(monkeypatch, isolated_settings, tmp_path):
    client, _ = _client(
        monkeypatch,
        isolated_settings,
        tmp_path,
        powerpoint_edit_enabled=True,
    )
    with client:
        thread_id = _start_thread(client)
        response = client.post(
            f"/chat/{thread_id}/upload",
            files={"files": ("report.pptx", _pptx_bytes(), "application/octet-stream")},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["errors"] == []
        assert body["files"][0]["filename"] == "report.pptx"


def test_upload_unknown_thread_returns_404(monkeypatch, isolated_settings, tmp_path):
    client, _ = _client(monkeypatch, isolated_settings, tmp_path)
    with client:
        response = client.post(
            "/chat/missing-thread/upload",
            files={"files": ("notes.txt", b"data", "text/plain")},
        )

        assert response.status_code == 404


def test_delete_thread_removes_uploaded_files(monkeypatch, isolated_settings, tmp_path):
    client, _ = _client(monkeypatch, isolated_settings, tmp_path)
    with client:
        thread_id = _start_thread(client)
        client.post(
            f"/chat/{thread_id}/upload",
            files={"files": ("notes.txt", b"hello", "text/plain")},
        )
        upload_dir = tmp_path / "files" / "chat_uploads" / thread_id
        assert upload_dir.exists()

        deleted = client.delete(f"/chat/{thread_id}")
        assert deleted.status_code == 200
        assert not upload_dir.exists()


def test_upload_edit_artifact_download_and_cleanup(
    monkeypatch,
    isolated_settings,
    tmp_path,
):
    docx = pytest.importorskip("docx")
    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("Original title")
    document.save(buffer)

    client, settings = _client(
        monkeypatch,
        isolated_settings,
        tmp_path,
        word_edit_enabled=True,
    )
    with client:
        thread_id = _start_thread(client)
        upload = client.post(
            f"/chat/{thread_id}/upload",
            files={
                "files": (
                    "report.docx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert upload.status_code == 200
        saved = upload.json()["files"][0]

        session_root = chat_api.session_upload_dir(settings, thread_id)
        edit = edit_word_document(
            saved["relative_path"],
            operations=[
                WordEditOperation(
                    action="replace_paragraph",
                    paragraph_index=0,
                    expected_text="Original title",
                    new_text="Updated title",
                )
            ],
            session_root=session_root,
            file_root=Path(settings.file_read_root),
            max_bytes=settings.file_read_max_bytes,
            thread_id=thread_id,
        )
        assert edit.artifact is not None, edit.content
        artifact = normalize_file_artifact(edit.artifact)

        assert artifact is not None
        download = client.get(artifact["url"])
        assert download.status_code == 200
        assert download.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        downloaded = docx.Document(io.BytesIO(download.content))
        assert downloaded.paragraphs[0].text == "Updated title"

        second_thread = _start_thread(client)
        assert (
            client.get(f"/chat/{second_thread}/files/{artifact['filename']}").status_code
            == 404
        )
        assert client.get(f"/chat/{thread_id}/files/missing.docx").status_code == 404
        assert client.get(f"/chat/unknown/files/{artifact['filename']}").status_code == 404

        deleted = client.delete(f"/chat/{thread_id}")
        assert deleted.status_code == 200
        assert not session_root.exists()


def test_upload_edit_excel_artifact_download_and_cleanup(
    monkeypatch,
    isolated_settings,
    tmp_path,
):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook, load_workbook

    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Item", "Qty"])
    sheet.append(["Widget", 10])
    workbook.save(buffer)

    client, settings = _client(
        monkeypatch,
        isolated_settings,
        tmp_path,
        excel_edit_enabled=True,
    )
    with client:
        thread_id = _start_thread(client)
        upload = client.post(
            f"/chat/{thread_id}/upload",
            files={
                "files": (
                    "report.xlsx",
                    buffer.getvalue(),
                    XLSX_MIME_TYPE,
                )
            },
        )
        assert upload.status_code == 200
        saved = upload.json()["files"][0]

        session_root = chat_api.session_upload_dir(settings, thread_id)
        edit = edit_excel(
            saved["relative_path"],
            operations=[
                ExcelEditOperation(
                    action="set_cell",
                    sheet="Data",
                    cell="A2",
                    expected_value="Widget",
                    value="Gizmo",
                )
            ],
            session_root=session_root,
            file_root=Path(settings.file_read_root),
            max_bytes=settings.file_read_max_bytes,
            thread_id=thread_id,
        )
        assert edit.artifact is not None, edit.content
        artifact = normalize_file_artifact(edit.artifact)

        assert artifact is not None
        assert artifact["mimeType"] == XLSX_MIME_TYPE
        download = client.get(artifact["url"])
        assert download.status_code == 200
        assert download.headers["content-type"].startswith(XLSX_MIME_TYPE)
        edited = load_workbook(io.BytesIO(download.content), data_only=False)
        assert edited["Data"]["A2"].value == "Gizmo"

        deleted = client.delete(f"/chat/{thread_id}")
        assert deleted.status_code == 200
        assert not session_root.exists()


def test_upload_edit_text_artifact_download_and_cleanup(
    monkeypatch,
    isolated_settings,
    tmp_path,
):
    client, settings = _client(
        monkeypatch,
        isolated_settings,
        tmp_path,
        text_edit_enabled=True,
    )
    with client:
        thread_id = _start_thread(client)
        upload = client.post(
            f"/chat/{thread_id}/upload",
            files={"files": ("notes.txt", b"alpha\r\nbeta\r\n", "text/plain")},
        )
        assert upload.status_code == 200
        saved = upload.json()["files"][0]

        session_root = chat_api.session_upload_dir(settings, thread_id)
        edit = edit_text_file(
            saved["relative_path"],
            operations=[
                TextEditOperation(
                    action="replace_line",
                    line_index=1,
                    expected_text="beta",
                    new_text="gamma",
                )
            ],
            session_root=session_root,
            file_root=Path(settings.file_read_root),
            max_bytes=settings.file_read_max_bytes,
            thread_id=thread_id,
        )
        assert edit.artifact is not None, edit.content
        artifact = normalize_file_artifact(edit.artifact)

        assert artifact is not None
        assert artifact["mimeType"] == "text/plain"
        download = client.get(artifact["url"])
        assert download.status_code == 200
        assert download.headers["content-type"].startswith("text/plain")
        assert download.content == b"alpha\r\ngamma\r\n"


def test_create_text_artifact_download_and_cleanup(
    monkeypatch,
    isolated_settings,
    tmp_path,
):
    client, settings = _client(
        monkeypatch,
        isolated_settings,
        tmp_path,
        text_edit_enabled=True,
    )
    with client:
        thread_id = _start_thread(client)
        session_root = chat_api.session_upload_dir(settings, thread_id)
        created = create_text_file(
            "todo",
            "first\nsecond\n",
            session_root=session_root,
            file_root=Path(settings.file_read_root),
            max_bytes=settings.file_read_max_bytes,
            thread_id=thread_id,
        )
        assert created.artifact is not None, created.content
        artifact = normalize_file_artifact(created.artifact)
        assert artifact is not None
        assert artifact["filename"] == "todo.txt"
        download = client.get(artifact["url"])
        assert download.status_code == 200
        assert download.content == b"first\nsecond\n"
        assert client.delete(f"/chat/{thread_id}").status_code == 200
        assert not session_root.exists()
