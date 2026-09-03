from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from pydantic import ValidationError

from src.backend.tools.excel_create import (
    XLSX_MIME_TYPE,
    ColumnFormat,
    SpreadsheetFormula,
    SpreadsheetSheet,
    build_excel_create_tools,
    create_excel_spreadsheet,
)


@pytest.fixture
def artifact_runtime() -> tuple[str, str]:
    node = shutil.which("node")
    modules = Path("node_modules").resolve()
    if node is None or not (modules / "@oai" / "artifact-tool").is_dir():
        pytest.skip("artifact-tool Node runtime is unavailable")
    return node, str(modules)


def sample_sheets() -> list[SpreadsheetSheet]:
    return [
        SpreadsheetSheet(
            name="Sales",
            rows=[
                ["Item", "Units", "Price", "Revenue", "Active"],
                ["Alpha", 2, 12.5, None, True],
                ["Beta", 3, 8.0, None, False],
            ],
            formulas=[
                SpreadsheetFormula(cell="D2", formula="=B2*C2", number_format="$#,##0.00"),
                SpreadsheetFormula(cell="D3", formula="=B3*C3", number_format="$#,##0.00"),
            ],
            number_formats=[ColumnFormat(column=2, format="$#,##0.00")],
            column_widths=[18, 12, 14, 16, 12],
        ),
        SpreadsheetSheet(
            name="Notes",
            rows=[["Status", "Comment"], ["Ready", "Verified workbook"]],
        ),
    ]


def test_create_excel_spreadsheet_is_session_scoped_and_collision_safe(
    tmp_path: Path,
    artifact_runtime: tuple[str, str],
) -> None:
    file_root = tmp_path / "files"
    session_root = file_root / "chat_uploads" / "thread-a"
    node, modules = artifact_runtime

    first = create_excel_spreadsheet(
        "quarterly report.xlsx",
        sheets=sample_sheets(),
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
        node_executable=node,
        node_modules_path=modules,
    )
    second = create_excel_spreadsheet(
        "quarterly report.xlsx",
        sheets=sample_sheets(),
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
        node_executable=node,
        node_modules_path=modules,
    )

    assert first.artifact is not None, first.content
    assert second.artifact is not None, second.content
    assert first.artifact["filename"] == "quarterly_report.xlsx"
    assert second.artifact["filename"] == "quarterly_report-2.xlsx"
    assert first.artifact["mimeType"] == XLSX_MIME_TYPE
    assert first.artifact["url"] == "/chat/thread-a/files/quarterly_report.xlsx"

    from openpyxl import load_workbook

    workbook = load_workbook(session_root / "quarterly_report.xlsx", data_only=False)
    assert workbook.sheetnames == ["Sales", "Notes"]
    assert workbook["Sales"]["A2"].value == "Alpha"
    assert workbook["Sales"]["B2"].value == 2
    assert workbook["Sales"]["E2"].value is True
    assert workbook["Sales"]["D2"].value == "=B2*C2"
    assert workbook["Sales"]["A1"].font.bold is True


def test_create_excel_spreadsheet_refuses_oversize_and_outside_session(
    tmp_path: Path,
    artifact_runtime: tuple[str, str],
) -> None:
    file_root = tmp_path / "files"
    session_root = file_root / "chat_uploads" / "thread-a"
    node, modules = artifact_runtime

    oversized = create_excel_spreadsheet(
        "too-large.xlsx",
        sheets=sample_sheets(),
        session_root=session_root,
        file_root=file_root,
        max_bytes=1,
        thread_id="thread-a",
        node_executable=node,
        node_modules_path=modules,
    )
    outside = create_excel_spreadsheet(
        "outside.xlsx",
        sheets=sample_sheets(),
        session_root=tmp_path / "elsewhere",
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
        node_executable=node,
        node_modules_path=modules,
    )

    assert oversized.artifact is None
    assert "too large" in oversized.content
    assert outside.artifact is None
    assert "outside the configured file root" in outside.content
    assert not list(tmp_path.rglob("*.xlsx"))


@pytest.mark.parametrize(
    "sheet",
    [
        {"name": "Bad/Name", "rows": [["A"]]},
        {"name": "Data", "rows": [["A", "B"], [1]]},
        {"name": "Data", "rows": [["A"]], "formulas": [{"cell": "B2", "formula": "=A1"}]},
        {"name": "Data", "rows": [["A"]], "formulas": [{"cell": "A1", "formula": "SUM(A1)"}]},
        {"name": "Data", "rows": [["bad\u0001value"]]},
    ],
)
def test_sheet_validation_rejects_unsafe_or_inconsistent_inputs(sheet: dict[str, object]) -> None:
    with pytest.raises(ValidationError):
        SpreadsheetSheet.model_validate(sheet)


def test_workbook_validation_rejects_duplicate_sheet_names() -> None:
    tool_schema = build_excel_create_tools
    assert callable(tool_schema)
    from src.backend.tools.excel_create import SpreadsheetCreateInput

    with pytest.raises(ValidationError):
        SpreadsheetCreateInput(
            filename="book.xlsx",
            sheets=[
                SpreadsheetSheet(name="Data", rows=[["A"]]),
                SpreadsheetSheet(name="data", rows=[["B"]]),
            ],
        )


def test_build_excel_create_tools_gates_and_supports_structured_invocation(
    isolated_settings,
    tmp_path: Path,
    artifact_runtime: tuple[str, str],
) -> None:
    node, modules = artifact_runtime
    session_root = tmp_path / "chat_uploads" / "thread-a"
    disabled = isolated_settings(file_read_enabled=True, excel_create_enabled=False)
    enabled = isolated_settings(
        file_read_enabled=True,
        excel_create_enabled=True,
        file_read_root=str(tmp_path),
        excel_node_executable=node,
        excel_node_modules_path=modules,
    )

    assert build_excel_create_tools(disabled, session_root=session_root, thread_id="thread-a") == []
    assert build_excel_create_tools(enabled, session_root=session_root, thread_id="../bad") == []
    tool = build_excel_create_tools(enabled, session_root=session_root, thread_id="thread-a")[0]
    message = tool.invoke(
        {
            "filename": "simple.xlsx",
            "sheets": [{"name": "Data", "rows": [["Name", "Score"], ["Alice", 90]]}],
        }
    )

    assert tool.name == "create_excel_spreadsheet"
    assert "Created simple.xlsx" in message
    assert (session_root / "simple.xlsx").is_file()
