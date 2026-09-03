from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from src.backend.tools.excel_edit import (
    XLSX_MIME_TYPE,
    ExcelEditOperation,
    build_excel_edit_tools,
    edit_excel,
    inspect_excel,
)


def _write_source_xlsx(path: Path) -> None:
    """Create a small .xlsx the edit tool can read and modify."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Item", "Qty", "Price", "Total"])
    sheet.append(["Widget", 10, 5, "=B2*C2"])
    sheet.append(["Gadget", 3, 12.5, "=B3*C3"])
    workbook.save(path)


def _roots(tmp_path: Path) -> tuple[Path, Path]:
    file_root = tmp_path / "files"
    session_root = file_root / "chat_uploads" / "thread-a"
    session_root.mkdir(parents=True, exist_ok=True)
    return file_root, session_root


def test_inspect_excel_lists_worksheets_cells_and_formulas(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    report = inspect_excel(
        "source.xlsx",
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
    )

    assert "WORKSHEETS (1):" in report
    assert "[0] Data" in report
    assert "[A1] Item" in report
    assert "[A2] Widget" in report
    assert "[B2] 10" in report
    assert "[D2] =B2*C2  (formula)" in report
    assert "[D3] =B3*C3  (formula)" in report


def test_edit_excel_applies_cell_formula_and_clear_operations(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    result = edit_excel(
        "source.xlsx",
        operations=[
            ExcelEditOperation(
                action="set_cell", sheet="Data", cell="A2",
                expected_value="Widget", value="Gizmo",
            ),
            ExcelEditOperation(
                action="set_formula", sheet="Data", cell="D2",
                expected_value="=B2*C2", formula="=B2*C2+1",
            ),
            ExcelEditOperation(
                action="clear_cell", sheet="Data", cell="A3",
                expected_value="Gadget",
            ),
        ],
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
    )

    assert result.artifact is not None, result.content
    assert result.artifact["mimeType"] == XLSX_MIME_TYPE
    assert result.artifact["filename"] == "source.edited.xlsx"
    assert result.artifact["url"] == "/chat/thread-a/files/source.edited.xlsx"

    edited = load_workbook(session_root / "source.edited.xlsx", data_only=False)
    assert edited["Data"]["A2"].value == "Gizmo"
    assert edited["Data"]["D2"].value == "=B2*C2+1"
    assert edited["Data"]["A3"].value is None

    # The original workbook is never modified.
    original = load_workbook(session_root / "source.xlsx", data_only=False)
    assert original["Data"]["A2"].value == "Widget"
    assert original["Data"]["D2"].value == "=B2*C2"
    assert original["Data"]["A3"].value == "Gadget"


def test_edit_excel_rejects_stale_expected_value_and_publishes_nothing(
    tmp_path: Path,
) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    result = edit_excel(
        "source.xlsx",
        operations=[
            ExcelEditOperation(
                action="set_cell", sheet="Data", cell="A2",
                expected_value="STALE", value="X",
            ),
        ],
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
    )

    assert result.artifact is None
    assert "re-inspect" in result.content
    assert not list(session_root.glob("source.edited*.xlsx"))


def test_edit_excel_adds_renames_and_deletes_worksheets(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    result = edit_excel(
        "source.xlsx",
        operations=[
            ExcelEditOperation(action="add_sheet", new_sheet_name="Summary"),
            ExcelEditOperation(action="rename_sheet", sheet="Data", new_sheet_name="Records"),
            ExcelEditOperation(action="delete_sheet", sheet="Summary"),
        ],
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
    )
    assert result.artifact is not None, result.content

    edited = load_workbook(session_root / "source.edited.xlsx")
    assert edited.sheetnames == ["Records"]


def test_edit_excel_refuses_to_delete_last_worksheet(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")  # only the "Data" sheet

    result = edit_excel(
        "source.xlsx",
        operations=[ExcelEditOperation(action="delete_sheet", sheet="Data")],
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
    )
    assert result.artifact is None
    assert "last worksheet" in result.content.lower()
    assert not list(session_root.glob("source.edited*.xlsx"))


def test_edit_excel_inserts_and_deletes_rows(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    result = edit_excel(
        "source.xlsx",
        operations=[
            ExcelEditOperation(action="insert_rows", sheet="Data", row_index=2, count=1),
            ExcelEditOperation(action="delete_rows", sheet="Data", row_index=4, count=1),
        ],
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
    )
    assert result.artifact is not None, result.content

    sheet = load_workbook(session_root / "source.edited.xlsx")["Data"]
    # Header intact, row 2 now blank (inserted), Widget shifted to row 3,
    # Gadget (formerly row 3, then row 4 after insert) deleted.
    assert sheet["A1"].value == "Item"
    assert sheet["A2"].value is None
    assert sheet["A3"].value == "Widget"


def test_edit_excel_inserts_and_deletes_columns(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    result = edit_excel(
        "source.xlsx",
        operations=[
            ExcelEditOperation(action="insert_columns", sheet="Data", column_index=4, count=1),
            ExcelEditOperation(action="delete_columns", sheet="Data", column_index=5, count=1),
        ],
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
    )
    assert result.artifact is not None, result.content

    sheet = load_workbook(session_root / "source.edited.xlsx")["Data"]
    # insert_columns(4) blanks column D; delete_columns(5) removes the old
    # Total column that shifted into E. Columns A-C stay put.
    assert sheet["A1"].value == "Item"
    assert sheet["C1"].value == "Price"
    assert sheet["D1"].value is None


def test_edit_excel_sets_column_width_and_number_format(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    result = edit_excel(
        "source.xlsx",
        operations=[
            ExcelEditOperation(action="set_column_width", sheet="Data", column_index=1, column_width=25),
            ExcelEditOperation(action="set_number_format", sheet="Data", column_index=2, number_format="#,##0.00"),
        ],
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
    )
    assert result.artifact is not None, result.content

    sheet = load_workbook(session_root / "source.edited.xlsx")["Data"]
    assert sheet.column_dimensions["A"].width == 25
    assert sheet.column_dimensions["B"].number_format == "#,##0.00"


def test_edit_excel_collision_safe_and_preserves_original(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")
    op = [ExcelEditOperation(action="set_cell", sheet="Data", cell="B2", expected_value=10, value=99)]

    first = edit_excel("source.xlsx", operations=op, session_root=session_root,
                       file_root=file_root, max_bytes=5_000_000, thread_id="thread-a")
    second = edit_excel("source.xlsx", operations=op, session_root=session_root,
                        file_root=file_root, max_bytes=5_000_000, thread_id="thread-a")

    assert first.artifact and first.artifact["filename"] == "source.edited.xlsx"
    assert second.artifact and second.artifact["filename"] == "source.edited-2.xlsx"

    first_value = load_workbook(session_root / "source.edited.xlsx", data_only=False)["Data"]["B2"].value
    second_value = load_workbook(session_root / "source.edited-2.xlsx", data_only=False)["Data"]["B2"].value
    assert first_value == 99
    assert second_value == 99
    # The uploaded source is never overwritten.
    assert load_workbook(session_root / "source.xlsx", data_only=False)["Data"]["B2"].value == 10


def test_edit_excel_refuses_access_outside_session(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    result = edit_excel(
        "../outside.xlsx",
        operations=[
            ExcelEditOperation(action="set_cell", sheet="Data", cell="A2",
                               expected_value="Widget", value="X"),
        ],
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id="thread-a",
    )
    assert result.artifact is None
    assert "access denied" in result.content.lower()


def test_edit_excel_refuses_oversize_source(tmp_path: Path) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    result = edit_excel(
        "source.xlsx",
        operations=[
            ExcelEditOperation(action="set_cell", sheet="Data", cell="A2",
                               expected_value="Widget", value="X"),
        ],
        session_root=session_root,
        file_root=file_root,
        max_bytes=1,
        thread_id="thread-a",
    )
    assert result.artifact is None
    assert "could not edit" in result.content.lower()


@pytest.mark.parametrize(
    "op",
    [
        {"action": "set_cell", "sheet": "Data", "cell": "A2", "expected_value": "x"},  # missing value
        {"action": "set_cell", "sheet": "Data", "cell": "A2", "value": "=B2"},  # value starts with '='
        {"action": "set_formula", "sheet": "Data", "cell": "A2", "formula": "B2"},  # no '='
        {"action": "set_cell", "sheet": "Data", "cell": "A0", "value": 1},  # bad cell address
        {"action": "set_cell", "cell": "A2", "value": 1},  # missing sheet
        {"action": "set_cell", "sheet": "Data", "value": 1},  # missing cell
        {"action": "insert_rows", "sheet": "Data"},  # missing row_index
        {"action": "rename_sheet", "sheet": "Data"},  # missing new_sheet_name
        {"action": "add_sheet"},  # missing new_sheet_name
        {"action": "set_column_width", "sheet": "Data", "column_index": 1},  # missing width
        {"action": "set_number_format", "sheet": "Data", "column_index": 1},  # missing format
        {"action": "delete_rows", "sheet": "Data", "row_index": 0},  # row_index < 1
        {"action": "set_formula", "sheet": "Data", "cell": "A2", "formula": "=cmd|'/c calc'!A1"},  # DDE blocked
        {"action": "bogus", "sheet": "Data"},  # invalid action literal
    ],
)
def test_edit_operation_validation_rejects_bad_inputs(op: dict) -> None:
    with pytest.raises(ValidationError):
        ExcelEditOperation.model_validate(op)


def test_build_excel_edit_tools_gates_and_supports_structured_invocation(
    isolated_settings,
    tmp_path: Path,
) -> None:
    file_root, session_root = _roots(tmp_path)
    _write_source_xlsx(session_root / "source.xlsx")

    disabled = isolated_settings(file_read_enabled=True, excel_edit_enabled=False)
    enabled = isolated_settings(
        file_read_enabled=True,
        excel_edit_enabled=True,
        file_read_root=str(tmp_path),
    )

    assert build_excel_edit_tools(disabled, session_root=session_root, thread_id="thread-a") == []
    assert build_excel_edit_tools(enabled, session_root=session_root, thread_id="../bad") == []

    inspect_tool, edit_tool = build_excel_edit_tools(
        enabled, session_root=session_root, thread_id="thread-a"
    )
    assert inspect_tool.name == "inspect_excel_spreadsheet"
    assert edit_tool.name == "edit_excel_spreadsheet"

    report = inspect_tool.invoke({"path": "source.xlsx"})
    assert "[A2] Widget" in report

    message = edit_tool.invoke(
        {
            "path": "source.xlsx",
            "operations": [
                {"action": "set_cell", "sheet": "Data", "cell": "A2",
                 "expected_value": "Widget", "value": "Gizmo"},
            ],
        }
    )
    assert "Applied" in message
    assert (session_root / "source.edited.xlsx").is_file()
    assert load_workbook(session_root / "source.edited.xlsx", data_only=False)["Data"]["A2"].value == "Gizmo"
