from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from pptx.dml.color import RGBColor
from pptx.util import Pt

from src.config import Settings
from src.graph.artifacts import normalize_artifact
from src.tools.powerpoint_edit import (
    PPTX_MIME_TYPE,
    PowerPointEditOperation,
    build_powerpoint_edit_tools,
    edit_powerpoint,
    inspect_powerpoint,
)

pptx = pytest.importorskip("pptx")


@pytest.fixture
def presentation_scope(tmp_path: Path) -> tuple[Path, Path, Path, str]:
    file_root = tmp_path / "files"
    thread_id = "thread-a"
    session_root = file_root / "chat_uploads" / thread_id
    session_root.mkdir(parents=True)
    source = session_root / "report.pptx"

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(0, 0, 4_000_000, 500_000)
    title_run = title.text_frame.paragraphs[0].add_run()
    title_run.text = "Quarterly report"
    title_run.font.bold = True
    title_run.font.size = Pt(32)
    title_run.font.color.rgb = RGBColor(255, 0, 0)
    table = slide.shapes.add_table(2, 2, 0, 600_000, 4_000_000, 1_000_000).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "East"
    table.cell(1, 1).text = "$120"
    revenue_run = table.cell(1, 1).text_frame.paragraphs[0].runs[0]
    revenue_run.font.bold = True
    revenue_run.font.color.rgb = RGBColor(0, 0, 255)
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, 4_000_000, 500_000).text = "Grouped text"
    presentation.save(source)
    return file_root, session_root, source, thread_id


def _edit(
    source: Path,
    *,
    file_root: Path,
    session_root: Path,
    thread_id: str,
    operations,
    output_name: str | None = None,
):
    return edit_powerpoint(
        source.name,
        operations=operations,
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
        thread_id=thread_id,
        output_name=output_name,
    )


def test_inspect_and_edit_text_and_table_cell(presentation_scope):
    file_root, session_root, source, thread_id = presentation_scope

    listing = inspect_powerpoint(
        source.name, session_root=session_root, file_root=file_root, max_bytes=5_000_000
    )
    assert "[0]" in listing
    assert "Quarterly report" in listing
    assert "[table row 1 col 1] $120" in listing
    assert "[2/0]" in listing
    assert "Grouped text" in listing

    result = _edit(
        source,
        file_root=file_root,
        session_root=session_root,
        thread_id=thread_id,
        operations=[
            PowerPointEditOperation(
                action="replace_text",
                slide_index=0,
                shape_path=[0],
                expected_text="Quarterly report",
                new_text="Annual report",
            ),
            PowerPointEditOperation(
                action="replace_table_cell",
                slide_index=0,
                shape_path=[1],
                row_index=1,
                column_index=1,
                expected_text="$120",
                new_text="$180",
            ),
            PowerPointEditOperation(
                action="replace_text",
                slide_index=0,
                shape_path=[2, 0],
                expected_text="Grouped text",
                new_text="Updated group",
            ),
        ],
    )

    assert "original file is unchanged" in result.content
    assert source.exists()
    output = session_root / "report.edited.pptx"
    assert output.exists()
    edited = pptx.Presentation(output)
    edited_title = edited.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert edited_title.text == "Annual report"
    assert edited_title.font.bold is True
    assert edited_title.font.size == Pt(32)
    assert edited_title.font.color.rgb == RGBColor(255, 0, 0)
    edited_revenue = edited.slides[0].shapes[1].table.cell(1, 1)
    assert edited_revenue.text == "$180"
    edited_revenue_run = edited_revenue.text_frame.paragraphs[0].runs[0]
    assert edited_revenue_run.font.bold is True
    assert edited_revenue_run.font.color.rgb == RGBColor(0, 0, 255)
    assert edited.slides[0].shapes[2].shapes[0].text == "Updated group"
    assert result.artifact is not None
    assert result.artifact["mimeType"] == PPTX_MIME_TYPE
    assert normalize_artifact(result.artifact) is not None


def test_edit_rejects_stale_expected_text_without_writing(presentation_scope):
    file_root, session_root, source, thread_id = presentation_scope
    result = _edit(
        source,
        file_root=file_root,
        session_root=session_root,
        thread_id=thread_id,
        operations=[
            PowerPointEditOperation(
                action="replace_text",
                slide_index=0,
                shape_path=[0],
                expected_text="Incorrect",
                new_text="Annual report",
            )
        ],
    )

    assert result.artifact is None
    assert "expected_text does not match" in result.content
    assert not (session_root / "report.edited.pptx").exists()


def test_inspect_wraps_malformed_presentation_errors(presentation_scope):
    file_root, session_root, _source, _thread_id = presentation_scope
    malformed = session_root / "malformed.pptx"
    with zipfile.ZipFile(malformed, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types>")
        archive.writestr("ppt/presentation.xml", "<p:presentation>")

    result = inspect_powerpoint(
        malformed.name,
        session_root=session_root,
        file_root=file_root,
        max_bytes=5_000_000,
    )

    assert result.startswith("Could not inspect PowerPoint presentation:")
    assert "invalid or unreadable" in result


def test_long_output_name_is_bounded_and_downloadable(presentation_scope):
    file_root, session_root, source, thread_id = presentation_scope
    result = _edit(
        source,
        file_root=file_root,
        session_root=session_root,
        thread_id=thread_id,
        output_name="a" * 200,
        operations=[
            PowerPointEditOperation(
                action="replace_text",
                slide_index=0,
                shape_path=[0],
                expected_text="Quarterly report",
                new_text="Annual report",
            )
        ],
    )

    assert result.artifact is not None, result.content
    assert len(result.artifact["filename"]) <= 200
    assert normalize_artifact(result.artifact) is not None


def test_tool_factory_is_gated_and_session_scoped(tmp_path: Path):
    settings = Settings(dashscope_api_key="test-key", file_read_enabled=True, powerpoint_edit_enabled=True)
    session_root = tmp_path / "chat_uploads" / "thread-a"
    assert build_powerpoint_edit_tools(settings) == []
    tools = build_powerpoint_edit_tools(settings, session_root=session_root, thread_id="thread-a")
    assert [tool.name for tool in tools] == ["inspect_powerpoint", "edit_powerpoint"]
